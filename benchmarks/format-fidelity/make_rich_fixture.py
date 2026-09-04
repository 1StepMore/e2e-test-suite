#!/usr/bin/env python3
"""Generate a deterministic "rich" DOCX fixture for the format-fidelity benchmark.

The fixture exercises the structural features the benchmark counts:
  * >= 1 table
  * >= 1 inline image (a tiny 1x1 px PNG, base64-embedded in this script, so
    the generator is fully self-contained — no network, no binary assets)
  * >= 1 built-in named paragraph style (Heading 1)
  * >= 1 custom paragraph style ("Benchmark Callout")

The script is idempotent: re-running it overwrites the same output file with
byte-identical content (fixed timestamps / revision ids are forced via
python-docx core properties), so the fixture is reproducible.

Usage:
    python make_rich_fixture.py                     # -> fixtures/rich.docx
    python make_rich_fixture.py [output.docx]      # explicit output path
"""

from __future__ import annotations

import base64
import sys
import tempfile
from pathlib import Path

from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.shared import Pt

# ---------------------------------------------------------------------------
# 1x1 px transparent PNG, base64-encoded inline (self-contained by design).
# ---------------------------------------------------------------------------
PNG_1X1_BASE64 = (
    "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk"
    "YPhfDwAChwGA60e6kgAAAABJRU5ErkJggg=="
)

OUTPUT_DEFAULT = Path(__file__).resolve().parent / "fixtures" / "rich.docx"


def build_document() -> Document:
    """Build the rich fixture document (deterministic content, no randomness)."""
    doc = Document()

    # Force deterministic core properties (python-docx defaults would embed a
    # creation timestamp that changes on every run).
    from datetime import datetime, timezone

    fixed_ts = datetime(2026, 9, 4, 0, 0, 0, tzinfo=timezone.utc)
    cp = doc.core_properties
    cp.created = fixed_ts
    cp.modified = fixed_ts
    cp.revision = 1
    cp.author = "format-fidelity benchmark"
    cp.title = "Rich fixture (benchmark)"

    # Custom paragraph style (>= 1 custom style requirement).
    callout = doc.styles.add_style("Benchmark Callout", WD_STYLE_TYPE.PARAGRAPH)
    callout.font.size = Pt(10)
    callout.font.italic = True

    # Content: built-in named style + body paragraphs + table + inline image.
    doc.add_paragraph("Rich Fixture", style="Heading 1")
    doc.add_paragraph(
        "This document exercises tables, inline images and named styles for "
        "the format-fidelity benchmark."
    )
    doc.add_paragraph("Callout paragraph with a custom style.", style="Benchmark Callout")

    table = doc.add_table(rows=2, cols=3)
    table.style = "Table Grid"
    header = table.rows[0].cells
    header[0].text = "Feature"
    header[1].text = "Count"
    header[2].text = "Notes"
    row = table.rows[1].cells
    row[0].text = "Table"
    row[1].text = "1"
    row[2].text = "2x3 grid with header row"

    doc.add_paragraph("Inline image below (1x1 px PNG):")
    # Fixed temp filename: python-docx derives the embedded picture's internal
    # name from the file name, so a random tmp name would break determinism.
    tmp_dir = tempfile.mkdtemp(prefix="omni_bench_")
    png_path = Path(tmp_dir) / "fixture.png"
    png_path.write_bytes(base64.b64decode(PNG_1X1_BASE64))
    try:
        doc.add_picture(str(png_path), width=Pt(1), height=Pt(1))
    finally:
        png_path.unlink(missing_ok=True)
        Path(tmp_dir).rmdir()

    return doc


def _rezip_with_fixed_timestamps(src: Path, dst: Path) -> None:
    """Rewrite the zip with fixed per-entry timestamps so re-runs are byte-identical.

    python-docx (via Python's zipfile) stamps every entry with the wall clock,
    which makes otherwise deterministic content produce a different container
    on each run.
    """
    import zipfile

    fixed_time = (2026, 9, 4, 0, 0, 0)
    with zipfile.ZipFile(src) as zin, zipfile.ZipFile(
        dst, "w", compression=zipfile.ZIP_DEFLATED
    ) as zout:
        for info in zin.infolist():
            zout.writestr(
                zipfile.ZipInfo(info.filename, date_time=fixed_time),
                zin.read(info.filename),
            )


def main() -> int:
    out_path = Path(sys.argv[1]) if len(sys.argv) > 1 else OUTPUT_DEFAULT
    doc = build_document()
    tmp_path = out_path.with_suffix(".tmp.docx")
    doc.save(str(tmp_path))
    _rezip_with_fixed_timestamps(tmp_path, out_path)
    tmp_path.unlink(missing_ok=True)
    print(f"Rich fixture written: {out_path}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
