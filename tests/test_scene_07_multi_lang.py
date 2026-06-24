"""Multi-language E2E test scenes (scene-07).

Verifies the OPP→OL→ORF pipeline runs without error for 3 language
pairs × 4 formats = 12 cells, using the FAKE_LLM seam. Does NOT
assert translation correctness (FAKE_LLM echoes input, not real
translation). Asserts: pipeline non-crash, output exists, output
is openable.

Requires: OMNI_TEST_FAKE_LLM=1, OMNI_TEST_FAKE_PANDOC=1
"""

from __future__ import annotations

import json
import os
import random
import subprocess
import zipfile
from pathlib import Path

import fitz
import pytest

random.seed(42)

# ---------------------------------------------------------------------------
# Shared content and helpers (also used by scene-07 conftest.py)
# ---------------------------------------------------------------------------

SCENE_DIR = Path(__file__).resolve().parent / "scenes" / "scene-07-multi-lang"
FIXTURES_DIR = SCENE_DIR / "fixtures"
EN_CONTENT = (FIXTURES_DIR / "en_content.md").read_text(encoding="utf-8")


def _make_docx(path: Path, content: str) -> None:
    """Generate a minimal .docx with headings and paragraphs."""
    from docx import Document

    doc = Document()
    for line in content.split("\n"):
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.strip():
            doc.add_paragraph(line)
    doc.save(str(path))


def _make_pdf(path: Path, content: str) -> None:
    """Generate a minimal .pdf via PyMuPDF (fitz)."""
    doc = fitz.open()
    page = doc.new_page()
    y = 50
    for line in content.split("\n"):
        if line.strip():
            page.insert_text((50, y), line)
            y += 20
    doc.save(str(path))
    doc.close()


# ---------------------------------------------------------------------------
# Test configuration
# ---------------------------------------------------------------------------

# Shared across all parametrized cells.
ALL_ENV = os.environ.copy()
ALL_ENV.update(
    {
        "OMNI_TEST_FAKE_LLM": "1",
        "OMNI_TEST_FAKE_PANDOC": "1",
        "OMNI_RATE_LIMIT_RPM": "0",
    }
)

LANG_PAIRS = ["en-fr", "en-ja", "en-ru"]
FORMATS = ["docx", "pdf", "html", "json"]

SUITE_ROOT = Path(__file__).resolve().parents[1]
OL_CONFIG = SUITE_ROOT / "Omni_Localizer" / "config" / "default.yaml"

# Build the 12-cell parameter list for indirect parametrization.
_SAMPLE_INPUT_PARAMS = [(lp, fmt) for lp in LANG_PAIRS for fmt in FORMATS]


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


@pytest.fixture(scope="session")
def sample_input(request, tmp_path_factory):
    """Generate an input file for one (lang_pair, fmt) cell.

    Uses ``indirect=True`` parametrization — ``request.param`` is a
    ``(lang_pair, fmt)`` tuple, e.g. ``("en-fr", "docx")``.

    Returns ``(path: Path, lang_pair: str, fmt: str)``.
    """
    lang_pair, fmt = request.param
    tmp = tmp_path_factory.mktemp(f"fixtures_{lang_pair}_{fmt}")

    if fmt == "docx":
        path = tmp / "sample.docx"
        _make_docx(path, EN_CONTENT)
    elif fmt == "pdf":
        path = tmp / "sample.pdf"
        _make_pdf(path, EN_CONTENT)
    elif fmt == "html":
        path = tmp / "sample.html"
        html_body = EN_CONTENT.replace("\n", "<br>\n")
        path.write_text(
            (
                "<!DOCTYPE html>\n<html><head><meta charset=\"utf-8\"></head>\n"
                f"<body><p>{html_body}</p></body>\n</html>\n"
            ),
            encoding="utf-8",
        )
    elif fmt == "json":
        path = tmp / "sample.json"
        data = {
            "company": "TechCorp",
            "name": "Alice Johnson",
            "email": "alice@techcorp.com",
            "phone": "+1-555-1234",
        }
        path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
    else:
        raise ValueError(f"Unsupported format: {fmt}")

    return (path, lang_pair, fmt)


# ---------------------------------------------------------------------------
# Tests
# ---------------------------------------------------------------------------


@pytest.mark.scene07
@pytest.mark.parametrize("sample_input", _SAMPLE_INPUT_PARAMS, indirect=True)
def test_pipeline_noncrash(sample_input, tmp_path: Path):
    """OPP→OL→ORF pipeline for a single scene-07 cell.

    ``sample_input`` is parametrized via ``indirect=True`` and returns
    ``(input_path, lang_pair, fmt)``.
    """
    input_path, lang_pair, fmt = sample_input
    src_lang, tgt_lang = lang_pair.split("-")

    # ------------------------------------------------------------------
    # Phase 1: OPP — extract to MD
    # ------------------------------------------------------------------
    opp_out = tmp_path / "opp"
    opp_out.mkdir()
    opp_cmd = [
        "opp",
        str(input_path),
        "--target-format", "md",
        "--source-lang", src_lang,
        "--target-lang", tgt_lang,
        "--output-dir", str(opp_out),
    ]
    opp_proc = subprocess.run(
        opp_cmd, capture_output=True, text=True, timeout=120, env=ALL_ENV,
    )
    assert opp_proc.returncode == 0, (
        f"OPP failed [{lang_pair}/{fmt}]:\nSTDOUT:\n{opp_proc.stdout[:2000]}\n"
        f"STDERR:\n{opp_proc.stderr[:2000]}"
    )

    md_files = list(opp_out.glob("*.md"))
    assert len(md_files) >= 1, f"No .md files in {opp_out} [{lang_pair}/{fmt}]"
    md_path = md_files[0]

    # ------------------------------------------------------------------
    # Phase 2: OL — translate MD
    # ------------------------------------------------------------------
    ol_out = tmp_path / "ol"
    ol_out.mkdir()
    ol_cmd = [
        "ol", "translate-md", str(md_path),
        "-c", str(OL_CONFIG),
        "-s", src_lang, "-t", tgt_lang,
        "-o", str(ol_out),
        "--no-cache",
    ]
    ol_proc = subprocess.run(
        ol_cmd, capture_output=True, text=True, timeout=120, env=ALL_ENV,
    )
    assert ol_proc.returncode == 0, (
        f"OL failed [{lang_pair}/{fmt}]:\nSTDOUT:\n{ol_proc.stdout[:2000]}\n"
        f"STDERR:\n{ol_proc.stderr[:2000]}"
    )

    ol_md_files = list(ol_out.glob("*.md"))
    assert len(ol_md_files) >= 1, (
        f"No translated .md in {ol_out} [{lang_pair}/{fmt}]"
    )
    ol_md = ol_md_files[0]

    # ------------------------------------------------------------------
    # Phase 3: ORF — backfill to target format
    # ------------------------------------------------------------------
    orf_out = tmp_path / "result"
    orf_out.mkdir()
    orf_cmd = [
        "orf", "apply-md", str(ol_md),
        "--target-format", fmt,
        "-o", str(orf_out / f"result.{fmt}"),
    ]
    orf_proc = subprocess.run(
        orf_cmd, capture_output=True, text=True, timeout=120, env=ALL_ENV,
    )
    assert orf_proc.returncode == 0, (
        f"ORF failed [{lang_pair}/{fmt}]:\nSTDOUT:\n{orf_proc.stdout[:2000]}\n"
        f"STDERR:\n{orf_proc.stderr[:2000]}"
    )

    # ------------------------------------------------------------------
    # Phase 4: Verify output
    # ------------------------------------------------------------------
    output_path = orf_out / f"result.{fmt}"
    assert output_path.exists(), (
        f"ORF output not found: {output_path} [{lang_pair}/{fmt}]"
    )
    assert output_path.stat().st_size > 0, (
        f"ORF output is empty: {output_path} [{lang_pair}/{fmt}]"
    )

    # Format-specific structure checks (confirm the file is "openable").
    if fmt == "json":
        data = json.loads(output_path.read_text(encoding="utf-8"))
        assert isinstance(data, (dict, list)), (
            f"JSON output is not a dict/list [{lang_pair}]"
        )
        _check_json_has_content(data, lang_pair)

    elif fmt == "html":
        content = output_path.read_text(encoding="utf-8").lower()
        assert "<html" in content or "<!doctype html" in content, (
            f"HTML output missing <html> tag [{lang_pair}]"
        )

    elif fmt == "docx":
        with zipfile.ZipFile(output_path) as zf:
            names = zf.namelist()
            assert "word/document.xml" in names, (
                f"DOCX missing word/document.xml [{lang_pair}]: {names}"
            )

    elif fmt == "pdf":
        # Minimal check: a valid PDF starts with "%PDF-"
        header = output_path.read_bytes()[:8]
        assert header.startswith(b"%PDF-"), (
            f"PDF output has wrong header [{lang_pair}]: {header!r}"
        )


def _check_json_has_content(data, lang_pair: str) -> None:
    """Walk JSON output and assert at least one non-empty string value."""
    def _walk(obj):
        if isinstance(obj, dict):
            for v in obj.values():
                yield from _walk(v)
        elif isinstance(obj, list):
            for v in obj:
                yield from _walk(v)
        else:
            yield str(obj)

    values = list(_walk(data))
    non_empty = [v for v in values if v.strip()]
    assert len(non_empty) > 0, (
        f"JSON output has no non-empty values [{lang_pair}]"
    )
