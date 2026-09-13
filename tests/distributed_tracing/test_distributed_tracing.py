"""Tests for Phase 4.5 cross-module distributed tracing.

Verifies W3C Trace Context propagation across the OPP → OL → ORF
pipeline:

- OPP's ``extract_document`` returns a ``traceparent`` field inside its
  ``content`` dict (W3C format ``00-{trace_id}-{span_id}-{flags}``)
  when ``OMNI_TRACING_ENABLED=1``.
- OL's ``translate_md_text`` accepts an optional ``traceparent``
  parameter; if provided, the OL span is a child of that trace
  and shares the same ``trace_id``.
- ORF's ``apply_md`` accepts an optional ``traceparent`` parameter;
  same child-span behavior.
- Parent-child relationship: OL span's ``parent_span_id`` equals
  OPP span's ``span_id``; ORF span's ``parent_span_id`` equals
  OL span's ``span_id``.

Hermetic: uses ``OMNI_TEST_FAKE_LLM=1`` and
``OMNI_TRACING_ENABLED=1``. No network, no real LLM, no
subprocesses — the test exercises the in-process tracing
modules and the in-process MCP tool functions directly.
"""
from __future__ import annotations

import importlib.util
import json
import re
import time
from pathlib import Path

import pytest

SUITE_ROOT = Path(__file__).resolve().parents[2]


# ---------------------------------------------------------------------------
# OTel reset + tracing loader
# ---------------------------------------------------------------------------


def _reset_otel_tracer_provider() -> None:
    """Reset the global OTel TracerProvider so a fresh one can
    be installed.  Mirrors the helper in tests/observability/
    tracing_health/test_tracing.py.
    """
    from opentelemetry import trace
    once = trace._TRACER_PROVIDER_SET_ONCE
    if once is not None:
        once._done = False
    trace._TRACER_PROVIDER = None


def _load_tracing_module(name: str, path: Path):
    """Load a tracing module fresh (per-test) and reset its
    setup state plus the OTel provider.
    """
    _reset_otel_tracer_provider()
    spec = importlib.util.spec_from_file_location(name, path)
    mod = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(mod)
    mod._setup_done = False
    return mod


def _read_jsonl(path: Path) -> list[dict]:
    if not path.exists():
        return []
    out: list[dict] = []
    with path.open("r", encoding="utf-8") as fh:
        for line in fh:
            line = line.strip()
            if not line:
                continue
            try:
                out.append(json.loads(line))
            except json.JSONDecodeError:
                continue
    return out


_TRACEPARENT_RE = re.compile(
    r"^[0-9a-f]{2}-([0-9a-f]{32})-([0-9a-f]{16})-([0-9a-f]{2})$"
)


def _parse_traceparent(tp: str) -> tuple[str, str, str]:
    """Parse a W3C traceparent string into (trace_id, span_id, flags)."""
    m = _TRACEPARENT_RE.match(tp.strip())
    assert m is not None, f"invalid traceparent: {tp!r}"
    return m.group(1), m.group(2), m.group(3)


# ---------------------------------------------------------------------------
# Fixtures
# ---------------------------------------------------------------------------


@pytest.fixture
def traces_dir(tmp_path, monkeypatch) -> Path:
    """Set up OTel env vars and point OMNI_TRACES_DIR at a tmp dir."""
    monkeypatch.setenv("OMNI_TRACING_ENABLED", "1")
    monkeypatch.setenv("OMNI_TRACES_DIR", str(tmp_path))
    monkeypatch.setenv("OMNI_TEST_FAKE_LLM", "1")
    monkeypatch.setenv("OMNI_TEST_FAKE_PANDOC", "1")
    monkeypatch.delenv("MCP_SHARED_SECRET", raising=False)
    # Point OL at the test config so FAKE_LLM works in end-to-end tests.
    ol_config = (
        SUITE_ROOT / "Omni_Localizer" / "config" / "test_universal.yaml"
    )
    if ol_config.exists():
        monkeypatch.setenv("OL_CONFIG_PATH", str(ol_config))
    return tmp_path


# ---------------------------------------------------------------------------
# Unit tests on the propagation helpers
# ---------------------------------------------------------------------------


class TestTraceparentHelpers:
    """The propagation helpers themselves: format, parse, round-trip."""

    def test_inject_traceparent_format(self, traces_dir):
        """``inject_traceparent`` returns a W3C-formatted string."""
        opp = _load_tracing_module(
            "opp_th_tracing",
            SUITE_ROOT / "Omni_Pre_Processor" / "src" / "opp" / "mcp" / "tracing.py",
        )
        with opp.start_call_tool_span("test", {}) as span:
            tp = opp.inject_traceparent(span)
        assert tp is not None
        trace_id, span_id, flags = _parse_traceparent(tp)
        assert len(trace_id) == 32
        assert len(span_id) == 16
        assert flags in ("00", "01")
        assert flags == "01", "new spans default to sampled"

    def test_extract_context_from_traceparent(self, traces_dir):
        """``_extract_traceparent_context`` parses a valid W3C header."""
        opp = _load_tracing_module(
            "opp_th_tracing_2",
            SUITE_ROOT / "Omni_Pre_Processor" / "src" / "opp" / "mcp" / "tracing.py",
        )
        ctx = opp._extract_traceparent_context(
            "00-aaaa1111aaaa1111aaaa1111aaaa1111-bbbb2222bbbb2222-01"
        )
        assert ctx is not None

    def test_extract_empty_string_returns_none(self, traces_dir):
        """Empty string is invalid, returns None."""
        opp = _load_tracing_module(
            "opp_th_tracing_3",
            SUITE_ROOT / "Omni_Pre_Processor" / "src" / "opp" / "mcp" / "tracing.py",
        )
        assert opp._extract_traceparent_context("") is None

    def test_extract_garbage_yields_empty_context(self, traces_dir):
        """Garbage string is invalid; OTel's W3C propagator returns
        an empty Context (semantics: "no valid parent"), not None.
        The helper passes this through unchanged so a span started
        with a garbage traceparent becomes a root span (no parent).
        """
        opp = _load_tracing_module(
            "opp_th_tracing_4",
            SUITE_ROOT / "Omni_Pre_Processor" / "src" / "opp" / "mcp" / "tracing.py",
        )
        ctx = opp._extract_traceparent_context("not-a-traceparent")
        # OTel semantics: invalid header → empty context (not None)
        assert ctx is not None
        from opentelemetry.context.context import Context
        assert isinstance(ctx, Context)

    def test_ol_inject_and_extract(self, traces_dir):
        """OL's helpers produce a valid W3C traceparent string."""
        ol = _load_tracing_module(
            "ol_th_tracing",
            SUITE_ROOT / "Omni_Localizer" / "src" / "ol_mcp" / "tracing.py",
        )
        with ol.start_call_tool_span("test", {}) as span:
            tp = ol.inject_traceparent(span)
        assert tp is not None
        trace_id, span_id, _ = _parse_traceparent(tp)
        assert trace_id != "0" * 32
        assert span_id != "0" * 16

    def test_orf_inject_and_extract(self, traces_dir):
        """ORF's helpers produce a valid W3C traceparent string."""
        orf = _load_tracing_module(
            "orf_th_tracing",
            SUITE_ROOT / "Omni_Re_Formatter" / "src" / "orf" / "mcp" / "tracing.py",
        )
        with orf.start_call_tool_span("test", {}) as span:
            tp = orf.inject_traceparent(span)
        assert tp is not None
        trace_id, span_id, _ = _parse_traceparent(tp)
        assert trace_id != "0" * 32
        assert span_id != "0" * 16


# ---------------------------------------------------------------------------
# Cross-module trace propagation — paired module tests
# ---------------------------------------------------------------------------


class TestDistributedTracePropagation:
    """End-to-end: paired-module tests that exercise the W3C trace
    context propagation.

    Each test resets OTel and loads the necessary modules fresh,
    so spans go to the correct JSONL file. This mirrors the
    real-world scenario where each MCP server is a separate
    process with its own OTel TracerProvider.
    """

    def test_opp_creates_root_span(self, traces_dir):
        """OPP creates a root span (no parent) with a valid trace_id."""
        opp = _load_tracing_module(
            "opp_dtp_root",
            SUITE_ROOT / "Omni_Pre_Processor" / "src" / "opp" / "mcp" / "tracing.py",
        )
        with opp.start_call_tool_span("extract_document", {}) as span:
            assert span is not None
            opp_tp = opp.inject_traceparent(span)
        assert opp_tp is not None
        opp_trace_id, opp_span_id, _ = _parse_traceparent(opp_tp)
        assert opp_trace_id != "0" * 32
        assert opp_span_id != "0" * 16

    def test_ol_span_inherits_opp_trace_id(self, traces_dir):
        """OL with ``traceparent`` from OPP produces a span whose
        ``trace_id`` matches OPP's, and ``parent_span_id`` equals
        OPP's ``span_id``.
        """
        # Step 1: OPP creates a root span
        opp = _load_tracing_module(
            "opp_dtp_ol",
            SUITE_ROOT / "Omni_Pre_Processor" / "src" / "opp" / "mcp" / "tracing.py",
        )
        with opp.start_call_tool_span("extract_document", {}) as opp_span:
            assert opp_span is not None
            opp_tp = opp.inject_traceparent(opp_span)
        opp_trace_id, opp_span_id, _ = _parse_traceparent(opp_tp)

        # Step 2: Reset OTel and load OL fresh
        ol = _load_tracing_module(
            "ol_dtp_ol",
            SUITE_ROOT / "Omni_Localizer" / "src" / "ol_mcp" / "tracing.py",
        )
        with ol.start_call_tool_span(
            "translate_md_text", {"content": "..."}, traceparent=opp_tp
        ) as ol_span:
            assert ol_span is not None
            ol_tp = ol.inject_traceparent(ol_span)
        ol_trace_id, ol_span_id, _ = _parse_traceparent(ol_tp)

        assert ol_trace_id == opp_trace_id, (
            f"OL trace_id {ol_trace_id} must equal OPP trace_id {opp_trace_id}"
        )
        assert ol_span_id != opp_span_id, "OL span_id must differ from OPP span_id"

        time.sleep(0.3)
        ol_spans = _read_jsonl(traces_dir / "ol.jsonl")
        assert len(ol_spans) >= 1, "OL spans should be written to ol.jsonl"
        last_ol = ol_spans[-1]
        assert last_ol["trace_id"] == opp_trace_id
        assert last_ol["parent_span_id"] == opp_span_id, (
            f"OL span's parent_span_id {last_ol['parent_span_id']} must equal "
            f"OPP span_id {opp_span_id}"
        )
        assert last_ol["name"] == "ol.call_tool"
        assert last_ol["attributes"]["module"] == "ol"

    def test_orf_span_inherits_ol_trace_id(self, traces_dir):
        """ORF with ``traceparent`` from OL produces a span whose
        ``trace_id`` matches OL's, and ``parent_span_id`` equals
        OL's ``span_id``.
        """
        # OL creates a root span
        ol = _load_tracing_module(
            "ol_dtp_orf",
            SUITE_ROOT / "Omni_Localizer" / "src" / "ol_mcp" / "tracing.py",
        )
        with ol.start_call_tool_span("translate_md_text", {}) as ol_span:
            ol_tp = ol.inject_traceparent(ol_span)
        ol_trace_id, ol_span_id, _ = _parse_traceparent(ol_tp)

        # ORF creates a child of OL
        orf = _load_tracing_module(
            "orf_dtp_orf",
            SUITE_ROOT / "Omni_Re_Formatter" / "src" / "orf" / "mcp" / "tracing.py",
        )
        with orf.start_call_tool_span(
            "apply_md", {"input_md": "/x.md"}, traceparent=ol_tp
        ) as orf_span:
            orf_tp = orf.inject_traceparent(orf_span)
        orf_trace_id, orf_span_id, _ = _parse_traceparent(orf_tp)

        assert orf_trace_id == ol_trace_id, (
            f"ORF trace_id {orf_trace_id} must equal OL trace_id {ol_trace_id}"
        )
        assert orf_span_id != ol_span_id

        time.sleep(0.3)
        orf_spans = _read_jsonl(traces_dir / "orf.jsonl")
        assert len(orf_spans) >= 1, "ORF spans should be written to orf.jsonl"
        last_orf = orf_spans[-1]
        assert last_orf["trace_id"] == ol_trace_id
        assert last_orf["parent_span_id"] == ol_span_id, (
            f"ORF span's parent_span_id {last_orf['parent_span_id']} must equal "
            f"OL span_id {ol_span_id}"
        )
        assert last_orf["name"] == "orf.call_tool"
        assert last_orf["attributes"]["module"] == "orf"

    def test_full_opp_ol_orf_chain(self, traces_dir):
        """Full OPP → OL → ORF chain: all 3 spans share the same
        trace_id; parent_span_id links them.
        """
        # Step 1: OPP root
        opp = _load_tracing_module(
            "opp_dtp_full",
            SUITE_ROOT / "Omni_Pre_Processor" / "src" / "opp" / "mcp" / "tracing.py",
        )
        with opp.start_call_tool_span("extract_document", {}) as opp_span:
            opp_tp = opp.inject_traceparent(opp_span)
        opp_trace_id, opp_span_id, _ = _parse_traceparent(opp_tp)

        # Step 2: OL child of OPP
        ol = _load_tracing_module(
            "ol_dtp_full",
            SUITE_ROOT / "Omni_Localizer" / "src" / "ol_mcp" / "tracing.py",
        )
        with ol.start_call_tool_span(
            "translate_md_text", {"content": "..."}, traceparent=opp_tp
        ) as ol_span:
            ol_tp = ol.inject_traceparent(ol_span)
        ol_trace_id, ol_span_id, _ = _parse_traceparent(ol_tp)
        assert ol_trace_id == opp_trace_id

        # Step 3: ORF child of OL
        orf = _load_tracing_module(
            "orf_dtp_full",
            SUITE_ROOT / "Omni_Re_Formatter" / "src" / "orf" / "mcp" / "tracing.py",
        )
        with orf.start_call_tool_span(
            "apply_md", {"input_md": "/x.md"}, traceparent=ol_tp
        ) as orf_span:
            orf_tp = orf.inject_traceparent(orf_span)
        orf_trace_id, orf_span_id, _ = _parse_traceparent(orf_tp)
        assert orf_trace_id == opp_trace_id

        # All 3 spans have different span_ids
        assert len({opp_span_id, ol_span_id, orf_span_id}) == 3, (
            "All 3 spans must have distinct span_ids"
        )

    def test_no_traceparent_creates_independent_traces(self, traces_dir):
        """When ``traceparent`` is not provided, each module
        starts its own (independent) trace.
        """
        opp = _load_tracing_module(
            "opp_dtp_indep",
            SUITE_ROOT / "Omni_Pre_Processor" / "src" / "opp" / "mcp" / "tracing.py",
        )
        with opp.start_call_tool_span("extract_document", {}) as opp_span:
            opp_tp = opp.inject_traceparent(opp_span)
        opp_trace_id, _, _ = _parse_traceparent(opp_tp)

        ol = _load_tracing_module(
            "ol_dtp_indep",
            SUITE_ROOT / "Omni_Localizer" / "src" / "ol_mcp" / "tracing.py",
        )
        with ol.start_call_tool_span("translate_md_text", {}) as ol_span:
            ol_tp = ol.inject_traceparent(ol_span)
        ol_trace_id, _, _ = _parse_traceparent(ol_tp)

        assert ol_trace_id != opp_trace_id, (
            "Without traceparent, OL must start a new trace (independent of OPP)"
        )

        time.sleep(0.3)
        ol_spans = _read_jsonl(traces_dir / "ol.jsonl")
        last_ol = ol_spans[-1]
        assert last_ol["parent_span_id"] is None, (
            "OL span without traceparent must have no parent_span_id (root span)"
        )


# ---------------------------------------------------------------------------
# End-to-end test: actual OPP/OL/ORF tool functions
# ---------------------------------------------------------------------------


class TestEndToEndDistributedTrace:
    """End-to-end test: calls the actual MCP tool functions and
    verifies the traceparent field in their responses.

    OPP's ``extract_document`` is called, then OL's
    ``translate_md_text`` is called with OPP's traceparent, then
    ORF's ``apply_md`` is called with OL's traceparent.  The
    spans are read from each module's JSONL file and the
    parent-child chain is verified.
    """

    def test_opp_ol_orf_share_single_trace(self, traces_dir, tmp_path):
        """Full pipeline: OPP → OL → ORF all in one trace."""
        from docx import Document  # type: ignore

        docx_path = tmp_path / "src.docx"
        doc = Document()
        doc.add_heading("Distributed Tracing", level=1)
        doc.add_paragraph(
            "This document is used to verify that the OPP extract_document, "
            "OL translate_md_text, and ORF apply_md tools can share a single "
            "OpenTelemetry trace across the pipeline via W3C trace context "
            "propagation."
        )
        doc.save(str(docx_path))

        from opp.mcp.config import MCPConfig  # type: ignore
        from opp.mcp.server import _init_server, _handle_call_tool  # type: ignore
        from ol_mcp.tools import _call_tool as _ol_call_tool  # type: ignore
        from orf.mcp.server import _call_tool as _orf_call_tool  # type: ignore
        import opp.mcp.tracing as _opp_tracing
        import ol_mcp.tracing as _ol_tracing
        import orf.mcp.tracing as _orf_tracing
        import asyncio

        def _reset_for_module(tracing_mod) -> None:
            """Reset OTel provider + the given tracing module's
            ``_setup_done`` flag so the next ``start_call_tool_span``
            call installs a fresh provider that writes to this
            module's JSONL file.
            """
            _reset_otel_tracer_provider()
            tracing_mod._setup_done = False

        opp_config = MCPConfig(
            allowed_directories=[str(tmp_path)],
            max_file_size_bytes=100 * 1024 * 1024,
            resource_storage_dir=str(tmp_path / "opp_resources"),
        )
        _init_server(opp_config)
        _reset_for_module(_opp_tracing)

        opp_resp_blocks = asyncio.run(
            _handle_call_tool(
                "extract_document",
                {
                    "file_path": str(docx_path),
                    "output_formats": ["md"],
                    "source_lang": "en",
                    "target_lang": "zh",
                },
            )
        )
        opp_payload = json.loads(opp_resp_blocks[0].text)
        assert opp_payload.get("success") is True, (
            f"OPP extract_document failed: {opp_payload}"
        )
        md_content = opp_payload["content"].get("md_content")
        assert md_content, "OPP must return md_content"
        opp_tp = opp_payload["content"].get("traceparent")
        assert opp_tp is not None, (
            "OPP extract_document must return traceparent when tracing is enabled"
        )
        opp_trace_id, opp_span_id, _ = _parse_traceparent(opp_tp)

        _reset_for_module(_ol_tracing)
        ol_resp_blocks = asyncio.run(
            _ol_call_tool(
                "translate_md_text",
                {
                    "content": md_content,
                    "source_lang": "en",
                    "target_lang": "zh",
                    "traceparent": opp_tp,
                },
            )
        )
        ol_payload = json.loads(ol_resp_blocks[0].text)
        assert ol_payload.get("success") is True, f"OL translate failed: {ol_payload}"
        ol_tp = ol_payload.get("traceparent")
        assert ol_tp is not None, (
            "OL translate_md_text must return traceparent when tracing is enabled"
        )
        ol_trace_id, ol_span_id, _ = _parse_traceparent(ol_tp)
        assert ol_trace_id == opp_trace_id, (
            f"OL trace_id {ol_trace_id} must equal OPP trace_id {opp_trace_id}"
        )

        translated_md = tmp_path / "translated.md"
        translated_md.write_text(ol_payload["translated"], encoding="utf-8")

        _reset_for_module(_orf_tracing)
        orf_out = tmp_path / "result.html"
        orf_resp_blocks = asyncio.run(
            _orf_call_tool(
                "apply_md",
                {
                    "input_md": str(translated_md),
                    "target_format": "html",
                    "output_path": str(orf_out),
                    "traceparent": ol_tp,
                },
            )
        )
        orf_payload = json.loads(orf_resp_blocks[0].text)
        orf_tp = orf_payload.get("traceparent")
        assert orf_tp is not None, (
            "ORF apply_md must return traceparent when tracing is enabled "
            f"(success={orf_payload.get('success')})"
        )
        orf_trace_id, orf_span_id, _ = _parse_traceparent(orf_tp)
        assert orf_trace_id == opp_trace_id, (
            f"ORF trace_id {orf_trace_id} must equal OPP trace_id {opp_trace_id}"
        )

        time.sleep(0.5)
        opp_spans = _read_jsonl(traces_dir / "opp.jsonl")
        ol_spans = _read_jsonl(traces_dir / "ol.jsonl")
        orf_spans = _read_jsonl(traces_dir / "orf.jsonl")

        opp_extract_spans = [s for s in opp_spans if s["attributes"].get("tool.name") == "extract_document"]
        assert len(opp_extract_spans) >= 1, "OPP should have emitted an extract_document span"
        opp_extract = opp_extract_spans[-1]
        assert opp_extract["trace_id"] == opp_trace_id
        assert opp_extract["parent_span_id"] is None, "OPP extract_document is the root span"

        ol_translate_spans = [s for s in ol_spans if s["attributes"].get("tool.name") == "translate_md_text"]
        assert len(ol_translate_spans) >= 1, "OL should have emitted a translate_md_text span"
        ol_translate = ol_translate_spans[-1]
        assert ol_translate["trace_id"] == opp_trace_id
        assert ol_translate["parent_span_id"] == opp_span_id, (
            f"OL translate_md_text parent_span_id {ol_translate['parent_span_id']} "
            f"must equal OPP span_id {opp_span_id}"
        )

        orf_apply_spans = [s for s in orf_spans if s["attributes"].get("tool.name") == "apply_md"]
        assert len(orf_apply_spans) >= 1, "ORF should have emitted an apply_md span"
        orf_apply = orf_apply_spans[-1]
        assert orf_apply["trace_id"] == opp_trace_id
        assert orf_apply["parent_span_id"] == ol_span_id, (
            f"ORF apply_md parent_span_id {orf_apply['parent_span_id']} "
            f"must equal OL span_id {ol_span_id}"
        )


# ---------------------------------------------------------------------------
# No-regression: existing 8 OTel tracing tests still pass (sanity check)
# ---------------------------------------------------------------------------


class TestNoRegression:
    def test_existing_tracing_module_apis_still_work(self, traces_dir):
        """Sanity: ``is_enabled()``, ``start_call_tool_span``,
        ``set_span_status``, and the JSONL exporter still work
        after the cross-module changes.
        """
        mod = _load_tracing_module(
            "opp_nr_tracing",
            SUITE_ROOT / "Omni_Pre_Processor" / "src" / "opp" / "mcp" / "tracing.py",
        )
        assert mod.is_enabled() is True
        with mod.start_call_tool_span("test", {}) as span:
            assert span is not None
            mod.set_span_status(span, "success", duration_ms=1.0)
        time.sleep(0.3)
