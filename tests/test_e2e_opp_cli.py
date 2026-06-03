"""E2E tests for OPP CLI (command-line interface).

Tests the OPP command-line tool invoked via subprocess:
- Test `opp --target-format=md`
- Test `opp --target-format=xlf --source-lang=en --target-lang=zh`
- Test `opp --target-format=both`
- Test batch mode `opp --batch`

Each test creates sample input files and verifies CLI output.
"""

import json
import subprocess
import sys
import tempfile
from pathlib import Path

import pytest


class TestOPPCLI:
    """OPP CLI end-to-end tests."""

    @pytest.mark.requires_opp
    def test_opp_target_format_md(self, tmp_path):
        """Test `opp --target-format=md` generates markdown output."""
        from docx import Document

        # Create sample DOCX
        doc = Document()
        doc.add_heading("Test Document", level=1)
        doc.add_paragraph("This is a test paragraph.")
        doc.add_paragraph("Another paragraph with content.")

        docx_path = tmp_path / "test_input.docx"
        doc.save(str(docx_path))

        output_dir = tmp_path / "output"
        output_dir.mkdir()

        # Run OPP CLI
        result = subprocess.run(
            [
                sys.executable, "-m", "opp",
                "--target-format=md",
                "--output-dir", str(output_dir),
                str(docx_path),
            ],
            capture_output=True,
            text=True,
        )

        assert result.returncode == 0, f"OPP CLI failed: {result.stderr}"

        # Check markdown was generated
        md_path = output_dir / "test_input.md"
        assert md_path.exists(), f"Markdown not generated: {md_path}"

        content = md_path.read_text(encoding="utf-8")
        assert "Test Document" in content
        assert "test paragraph" in content

    @pytest.mark.requires_opp
    def test_opp_target_format_xlf(self, tmp_path):
        """Test `opp --target-format=xlf --source-lang=en --target-lang=zh` generates XLIFF."""
        from docx import Document

        # Create sample DOCX
        doc = Document()
        doc.add_heading("Hello World", level=1)
        doc.add_paragraph("This is the content to translate.")

        docx_path = tmp_path / "hello.docx"
        doc.save(str(docx_path))

        output_dir = tmp_path / "output"
        output_dir.mkdir()

        # Run OPP CLI with XLIFF target format
        result = subprocess.run(
            [
                sys.executable, "-m", "opp",
                "--target-format=xlf",
                "--source-lang=en",
                "--target-lang=zh",
                "--output-dir", str(output_dir),
                str(docx_path),
            ],
            capture_output=True,
            text=True,
        )

        assert result.returncode == 0, f"OPP CLI XLIFF failed: {result.stderr}"

        # Check XLIFF was generated
        xlf_path = output_dir / "hello.xlf"
        assert xlf_path.exists(), f"XLIFF not generated: {xlf_path}"

        xliff_content = xlf_path.read_text(encoding="utf-8")
        assert "<xliff" in xliff_content or "<Xliff" in xliff_content
        assert "Hello World" in xliff_content or "Hello" in xliff_content

        # Verify manifest was also created
        manifest_path = output_dir / "hello_manifest.json"
        assert manifest_path.exists(), "Manifest not generated"

        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        assert manifest["tool"] == "OPP"
        assert manifest["extraction"]["source_lang"] == "en"
        assert manifest["extraction"]["target_lang"] == "zh"

    @pytest.mark.requires_opp
    def test_opp_target_format_both(self, tmp_path):
        """Test `opp --target-format=both` generates both MD and XLIFF."""
        from docx import Document

        # Create sample DOCX
        doc = Document()
        doc.add_heading("Sample Title", level=1)
        doc.add_paragraph("First paragraph content.")
        doc.add_paragraph("Second paragraph here.")

        docx_path = tmp_path / "sample.docx"
        doc.save(str(docx_path))

        output_dir = tmp_path / "output"
        output_dir.mkdir()

        # Run OPP CLI with both formats
        result = subprocess.run(
            [
                sys.executable, "-m", "opp",
                "--target-format=both",
                "--source-lang=en",
                "--target-lang=zh",
                "--output-dir", str(output_dir),
                str(docx_path),
            ],
            capture_output=True,
            text=True,
        )

        assert result.returncode == 0, f"OPP CLI both failed: {result.stderr}"

        # Check both MD and XLIFF were generated
        md_path = output_dir / "sample.md"
        xlf_path = output_dir / "sample.xlf"

        assert md_path.exists(), "Markdown not generated"
        assert xlf_path.exists(), "XLIFF not generated"

        md_content = md_path.read_text(encoding="utf-8")
        assert "Sample Title" in md_content

        xliff_content = xlf_path.read_text(encoding="utf-8")
        assert "<xliff" in xliff_content or "<Xliff" in xliff_content

        # Check manifest
        manifest_path = output_dir / "sample_manifest.json"
        manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
        assert manifest["extraction"]["outputs"]["markdown"]["path"] is not None
        assert manifest["extraction"]["outputs"]["xliff"]["path"] is not None

    @pytest.mark.requires_opp
    def test_opp_batch_mode(self, tmp_path):
        """Test batch mode `opp --batch` processes multiple files."""
        from docx import Document

        # Create multiple DOCX files
        for i in range(3):
            doc = Document()
            doc.add_heading(f"Document {i}", level=1)
            doc.add_paragraph(f"Content for document {i}.")

            docx_path = tmp_path / f"batch_{i}.docx"
            doc.save(str(docx_path))

        output_dir = tmp_path / "output"
        output_dir.mkdir()

        # Run OPP CLI in batch mode
        result = subprocess.run(
            [
                sys.executable, "-m", "opp",
                "--batch",
                "--target-format=md",
                "--output-dir", str(output_dir),
                str(tmp_path),
            ],
            capture_output=True,
            text=True,
        )

        assert result.returncode == 0, f"OPP batch failed: {result.stderr}"

        # Check all files were processed
        for i in range(3):
            md_path = output_dir / f"batch_{i}.md"
            assert md_path.exists(), f"Batch file {i} not processed: {md_path}"

            content = md_path.read_text(encoding="utf-8")
            assert f"Document {i}" in content

    @pytest.mark.requires_opp
    def test_opp_detect_format(self, tmp_path):
        """Test `opp --detect-format` auto-detects file format."""
        from docx import Document

        doc = Document()
        doc.add_heading("Format Test", level=1)
        doc.add_paragraph("Content here.")

        docx_path = tmp_path / "detect_test.docx"
        doc.save(str(docx_path))

        # Run OPP with --detect-format
        result = subprocess.run(
            [
                sys.executable, "-m", "opp",
                "--detect-format",
                str(docx_path),
            ],
            capture_output=True,
            text=True,
        )

        # Should complete without error
        assert result.returncode == 0, f"OPP detect-format failed: {result.stderr}"

    @pytest.mark.requires_opp
    def test_opp_verbose_output(self, tmp_path):
        """Test `opp -v` verbose output is produced."""
        from docx import Document

        doc = Document()
        doc.add_heading("Verbose Test", level=1)
        doc.add_paragraph("Content.")

        docx_path = tmp_path / "verbose.docx"
        doc.save(str(docx_path))

        output_dir = tmp_path / "output"
        output_dir.mkdir()

        result = subprocess.run(
            [
                sys.executable, "-m", "opp",
                "-v",
                "--target-format=md",
                "--output-dir", str(output_dir),
                str(docx_path),
            ],
            capture_output=True,
            text=True,
        )

        assert result.returncode == 0
        # Verbose output should contain processing info
        output = result.stdout + result.stderr
        assert len(output) > 0

    @pytest.mark.requires_opp
    def test_opp_no_files_error(self):
        """Test OPP CLI returns error when no files provided."""
        result = subprocess.run(
            [sys.executable, "-m", "opp"],
            capture_output=True,
            text=True,
        )

        assert result.returncode != 0

    @pytest.mark.requires_opp
    def test_opp_xlf_requires_target_lang(self, tmp_path):
        """Test OPP CLI requires --target-lang when using --target-format=xlf."""
        from docx import Document

        doc = Document()
        doc.add_heading("Title", level=1)
        doc.add_paragraph("Content.")

        docx_path = tmp_path / "test.docx"
        doc.save(str(docx_path))

        # Run without --target-lang (should fail)
        result = subprocess.run(
            [
                sys.executable, "-m", "opp",
                "--target-format=xlf",
                str(docx_path),
            ],
            capture_output=True,
            text=True,
        )

        assert result.returncode != 0
        assert "target-lang" in result.stderr or "--target-lang" in result.stderr