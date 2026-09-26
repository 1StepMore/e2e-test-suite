"""E2E tests for OPP MCP server tools — real shipped surface.

Tests the OPP MCP server tools:
- extract_document tool
- batch_extract tool
- detect_format tool

Each test calls the real MCP tool functions after initializing the real
server state through the shipped public entry points — ``load_config()``
(env-driven, fail-closed) + ``_init_server()`` — exactly as
``opp.mcp.server.main()`` does at startup. No server internals are
patched: since the tools were split into ``opp.mcp.tools.*``, they read
their shared state (``_config`` / ``_validator`` / ``_pipeline`` /
``_serializer``) from ``opp.mcp.common``. The names re-exported by
``opp.mcp.server`` are PEP 562 delegations, so patching
``opp.mcp.server._validator`` cannot affect tool behavior, and skipping
``_init_server`` fails closed with ``McpError("Server not initialized")``.

Response envelope contract (``opp.mcp._errors.mcp_error_boundary`` +
``opp.mcp.server._handle_call_tool``):
- success: ``{"success": True, "content": {...}}``
- error:   ``{"success": False, "error": {"code", "message"},
            "error_code", "message", "recovery": {"strategy", "hint"},
            **extra}`` (e.g. ``validation_errors`` for batch path denial)

Environment is explicit and deterministic: the MCP path is allowlist-gated
and fail-closed, so every test runs with ``OPP_MCP_ALLOWED_DIRS`` bound to
its own ``tmp_path`` and auth disabled.
"""

import asyncio
from dataclasses import replace
from pathlib import Path

import pytest


@pytest.fixture()
def opp_mcp_server(tmp_path, monkeypatch):
    """Initialize the REAL OPP MCP server state against ``tmp_path``.

    Mirrors ``opp.mcp.server.main()``: ``load_config()`` from environment,
    then ``_init_server(config)``. The allowlist is bound to ``tmp_path``
    so path validation is deterministic; resource output is redirected
    into ``tmp_path`` because the default ``./mcp_resources`` is
    CWD-relative.
    """
    from opp.mcp.common import _init_server
    from opp.mcp.config import load_config

    # Explicit, deterministic, fail-closed environment.
    monkeypatch.setenv("OPP_MCP_ALLOWED_DIRS", str(tmp_path))
    monkeypatch.delenv("MCP_ALLOWED_DIRECTORIES", raising=False)
    monkeypatch.delenv("MCP_SHARED_SECRET", raising=False)  # auth disabled

    config = replace(
        load_config(),
        resource_storage_dir=tmp_path / "mcp_resources",
    )
    _init_server(config)
    return config


def _make_docx(path: Path, heading: str, paragraph: str) -> None:
    """Create a minimal real DOCX with one heading and one paragraph."""
    from docx import Document

    doc = Document()
    doc.add_heading(heading, level=1)
    doc.add_paragraph(paragraph)
    doc.save(str(path))


class TestOPPMCP:
    """OPP MCP server tool tests against the real shipped surface."""

    @pytest.mark.requires_opp
    def test_extract_document_md_format(self, opp_mcp_server, tmp_path):
        """extract_document with output_formats=["md"] returns the real
        success envelope: md_content, suggested_pipeline, and skeleton_path
        nested under ``content``."""
        docx_path = tmp_path / "test.docx"
        _make_docx(docx_path, "Test Document", "Test paragraph content.")

        from opp.mcp.server import extract_document

        result = asyncio.run(
            extract_document(
                file_path=str(docx_path),
                output_formats=["md"],
                source_lang="en",
                target_lang="zh",
            )
        )

        assert result["success"] is True
        content = result["content"]
        assert "Test paragraph content." in content["md_content"]
        assert "Test Document" in content["md_content"]
        # DOCX supports both MD and XLIFF → suggested_pipeline is "both".
        assert content["suggested_pipeline"] == "both"
        # DOCX extraction always ships a skeleton.zip next to the input.
        assert "skeleton_path" in content
        assert Path(content["skeleton_path"]).exists()

    @pytest.mark.requires_opp
    def test_extract_document_xlf_format(self, opp_mcp_server, tmp_path):
        """extract_document with output_formats=["xlf"] returns real XLIFF
        content with a positive trans-unit count."""
        docx_path = tmp_path / "xliff_test.docx"
        _make_docx(docx_path, "XLIFF Test", "Content for translation.")

        from opp.mcp.server import extract_document

        result = asyncio.run(
            extract_document(
                file_path=str(docx_path),
                output_formats=["xlf"],
                source_lang="en",
                target_lang="zh",
            )
        )

        assert result["success"] is True
        content = result["content"]
        assert "<trans-unit" in content["xliff_content"]
        assert content["xliff_units_count"] >= 1
        # md generation must NOT run when only xlf was requested.
        assert "md_content" not in content
        assert "skeleton_path" in content

    @pytest.mark.requires_opp
    def test_batch_extract_validation_error(self, opp_mcp_server):
        """batch_extract fails fast on an invalid path with the real error
        envelope (OPP_PATH_DENIED + validation_errors + recovery)."""
        from opp.mcp.server import batch_extract

        bad_path = "/nonexistent/path/file.docx"
        result = asyncio.run(
            batch_extract(
                file_paths=[bad_path],
                output_formats=["md"],
            )
        )

        assert result["success"] is False
        assert result["error"]["code"] == "OPP_PATH_DENIED"
        assert result["error_code"] == "OPP_PATH_DENIED"
        assert result["validation_errors"][0]["file_path"] == bad_path
        assert result["recovery"]["strategy"] == "use_allowed_path"

    @pytest.mark.requires_opp
    def test_detect_format_tool(self, opp_mcp_server, tmp_path):
        """detect_format_tool identifies a DOCX via magic bytes, returning
        the format and confidence nested under ``content``."""
        docx_path = tmp_path / "detect.docx"
        _make_docx(docx_path, "Format Detection", "Content.")

        from opp.mcp.server import detect_format_tool

        result = asyncio.run(detect_format_tool(file_path=str(docx_path)))

        assert result["success"] is True
        assert result["content"]["format"] == "docx"  # FormatType.DOCX.value
        confidence = result["content"]["confidence"]
        assert isinstance(confidence, float)
        assert 0.0 <= confidence <= 1.0

    @pytest.mark.requires_opp
    def test_detect_format_tool_invalid_path(self, opp_mcp_server):
        """detect_format_tool rejects a path outside the allowlist with the
        real error envelope (OPP_PATH_DENIED + recovery hint)."""
        from opp.mcp.server import detect_format_tool

        result = asyncio.run(detect_format_tool(file_path="/invalid/path/file.docx"))

        assert result["success"] is False
        assert result["error"]["code"] == "OPP_PATH_DENIED"
        assert result["error_code"] == "OPP_PATH_DENIED"
        assert result["message"]
        assert result["recovery"]["strategy"] == "use_allowed_path"

    @pytest.mark.requires_opp
    def test_extract_document_both_formats(self, opp_mcp_server, tmp_path):
        """extract_document with output_formats=["both"] returns BOTH
        md_content and xliff_content in one envelope."""
        docx_path = tmp_path / "both.docx"
        _make_docx(docx_path, "Both Formats Test", "Testing MD and XLIFF output.")

        from opp.mcp.server import extract_document

        result = asyncio.run(
            extract_document(
                file_path=str(docx_path),
                output_formats=["both"],
                source_lang="en",
                target_lang="zh",
            )
        )

        assert result["success"] is True
        content = result["content"]
        assert "Testing MD and XLIFF output." in content["md_content"]
        assert "<trans-unit" in content["xliff_content"]
        assert content["suggested_pipeline"] == "both"
        assert "skeleton_path" in content

    @pytest.mark.requires_opp
    def test_batch_extract_mixed_results(self, opp_mcp_server, tmp_path):
        """batch_extract fails closed when ANY path is denied: the whole
        batch aborts with OPP_PATH_DENIED + validation_errors naming only
        the denied path — the valid file is NOT extracted.

        (The shipped surface validates ALL paths upfront and raises before
        any extraction runs, so partial per-file results can never occur
        for a batch containing a denied path.)
        """
        valid_path = tmp_path / "valid.docx"
        _make_docx(valid_path, "Mixed", "Content.")

        invalid_path = "/invalid/path.docx"

        from opp.mcp.server import batch_extract

        result = asyncio.run(
            batch_extract(
                file_paths=[str(valid_path), invalid_path],
                output_formats=["md"],
            )
        )

        assert result["success"] is False
        assert result["error"]["code"] == "OPP_PATH_DENIED"
        assert result["error_code"] == "OPP_PATH_DENIED"
        denied = {e["file_path"] for e in result["validation_errors"]}
        assert denied == {invalid_path}
        # Fail-closed: no partial extraction results were produced.
        assert "content" not in result

    @pytest.mark.requires_opp
    def test_extract_document_with_resource_dir(self, opp_mcp_server, tmp_path):
        """extract_document accepts a resource_dir inside the allowlist and
        returns the real success envelope."""
        docx_path = tmp_path / "resource.docx"
        _make_docx(docx_path, "Resource Test", "Content.")

        resource_dir = tmp_path / "resources"
        resource_dir.mkdir()

        from opp.mcp.server import extract_document

        result = asyncio.run(
            extract_document(
                file_path=str(docx_path),
                output_formats=["md"],
                resource_dir=str(resource_dir),
            )
        )

        assert result["success"] is True
        content = result["content"]
        assert "Content." in content["md_content"]
        assert content["suggested_pipeline"] == "both"
