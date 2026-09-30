"""OPP#80 Wave 3: backward compatibility for the ORF ``_para`` regex change.

The Wave 1 ORF change made the table-cell resname regex accept an OPTIONAL
``_para{p}`` suffix: ``^table_(\\d+)_r(\\d+)_c(\\d+)(?:_para(\\d+))?$``. The
safety property that matters for older XLIFF is: **a bare resname must take
exactly the legacy whole-cell path** — the new optional group must not perturb
the old behavior.

These suite-level tests drive the REAL shipped ORF CLI (no ORF internals, no
mocked loader) against a real DOCX skeleton, and pin the legacy outcome against
a golden value. They also confirm the documented, accepted failure mode of a
NEW ``_para`` XLIFF fed to the same writer: an out-of-range paragraph index is
skipped and the cell is left intact (the warning detail itself is unit-tested
in ORF's own ``tests/test_xliff2docx_table_cells.py``).

Hermetic: no LLM, no network, ``tmp_path`` only.
"""

from __future__ import annotations

import os
import subprocess
import sys
import xml.etree.ElementTree as ET
import zipfile
from pathlib import Path

import pytest

pytestmark = [pytest.mark.e2e, pytest.mark.requires_orf]

_SUITE_ROOT = Path(__file__).resolve().parent.parent
_SRC_DIRS = (
    _SUITE_ROOT / "Omni_Pre_Processor" / "src",
    _SUITE_ROOT / "Omni_Localizer" / "src",
    _SUITE_ROOT / "Omni_Re_Formatter" / "src",
)

W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def _subprocess_env() -> dict:
    env = os.environ.copy()
    parts = [str(d) for d in _SRC_DIRS if d.exists()]
    existing = env.get("PYTHONPATH", "")
    if existing:
        parts.append(existing)
    env["PYTHONPATH"] = ":".join(parts)
    env["OMNI_TEST_FAKE_PANDOC"] = "1"
    return env


def _build_two_para_cell_docx(path: Path) -> None:
    """A DOCX whose only table cell has two direct ``w:p``: ``A`` and ``B``."""
    from docx import Document

    doc = Document()
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    cell.text = "A"
    cell.add_paragraph("B")
    doc.save(str(path))


def _write_xliff(path: Path, unit_xml: str) -> None:
    path.write_text(
        '<?xml version="1.0" encoding="UTF-8"?>\n'
        '<xliff version="1.2" xmlns="urn:oasis:names:tc:xliff:document:1.2">\n'
        '  <file original="t" source-language="en" target-language="zh">\n'
        f'    <body>\n{unit_xml}\n    </body>\n  </file>\n</xliff>\n',
        encoding="utf-8",
    )


def _unit(uid: str, resname: str, source: str, target: str) -> str:
    return (
        f'      <trans-unit xml:space="preserve" id="{uid}" resname="{resname}">\n'
        f'        <source xml:space="preserve">{source}</source>\n'
        f'        <target xml:space="preserve">{target}</target>\n'
        "      </trans-unit>"
    )


def _run_orf(skeleton: Path, xliff: Path, output: Path) -> None:
    result = subprocess.run(
        [
            sys.executable, "-m", "orf.cli", "apply-xliff", str(skeleton),
            "--xliff", str(xliff), "--output", str(output), "--format", "docx",
        ],
        capture_output=True, text=True, env=_subprocess_env(), timeout=120,
    )
    assert result.returncode == 0, (
        f"ORF CLI failed (rc={result.returncode})\n"
        f"stdout tail: {result.stdout[-2000:]}\n"
        f"stderr tail: {result.stderr[-2000:]}"
    )
    assert output.exists()


def _cell_paragraphs(docx_path: Path) -> list[str]:
    with zipfile.ZipFile(docx_path) as zf:
        root = ET.fromstring(zf.read("word/document.xml"))
    cell = next(root.iter(f"{W_NS}tc"))
    return [
        "".join(t.text or "" for t in p.iter(f"{W_NS}t"))
        for p in cell.findall(f"{W_NS}p")
    ]


def test_bare_resname_uses_legacy_whole_cell_write(tmp_path: Path):
    """New ORF + OLD (bare) XLIFF must equal the legacy whole-cell behaviour.

    Golden value: the whole-cell write puts the target into the cell's FIRST
    paragraph and clears the tail runs, so a two-paragraph cell becomes
    ``["ZH-WHOLE", ""]``. This is the byte-for-byte legacy outcome the optional
    ``_para`` regex group must not change.
    """
    skeleton = tmp_path / "source.docx"
    _build_two_para_cell_docx(skeleton)
    xliff = tmp_path / "bare.xlf"
    _write_xliff(xliff, _unit("1", "table_0_r0_c0", "A\nB", "ZH-WHOLE"))
    output = tmp_path / "out.docx"

    _run_orf(skeleton, xliff, output)

    assert _cell_paragraphs(output) == ["ZH-WHOLE", ""]


def test_out_of_range_para_resname_skips_and_leaves_cell_intact(tmp_path: Path):
    """A NEW ``_para{p}`` XLIFF with an out-of-range index must not corrupt.

    Documented accepted failure mode: the positional write declines (the
    warning itself is covered by ORF's own unit test
    ``test_para_out_of_range_skips_and_warns``), the cell keeps its source text,
    and the output stays parseable. This is the suite-level integration of that
    property against a real skeleton.
    """
    skeleton = tmp_path / "source.docx"
    _build_two_para_cell_docx(skeleton)
    xliff = tmp_path / "oor.xlf"
    _write_xliff(
        xliff, _unit("1", "table_0_r0_c0_para9", "A", "ZH-PARA9")
    )
    output = tmp_path / "out.docx"

    _run_orf(skeleton, xliff, output)

    # Cell untouched; output is still a readable DOCX.
    assert _cell_paragraphs(output) == ["A", "B"]
