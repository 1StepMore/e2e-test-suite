"""E2E tests for ORF - All output formats.

Tests MD → DOCX, PPTX, EPUB, HTML, ODT, PDF, RTF, ICML, SRT, XLSX, CSV, JSON
and XLIFF → DOCX, PPTX, EPUB, HTML, ODF conversions.

Each test:
1. Creates input MD/XLIFF
2. Calls the appropriate ORF converter
3. Verifies output file exists and is valid
"""

import json
import sys
import tempfile
import zipfile
from io import BytesIO
from pathlib import Path

import pytest

# Ensure ORF src is in path
sys.path.insert(0, str(Path(__file__).parent.parent / "Omni_Re_Formatter" / "src"))


# =============================================================================
# Sample Fixtures
# =============================================================================


@pytest.fixture
def sample_md(tmp_path: Path) -> Path:
    """Create a sample markdown file for conversion tests."""
    content = """# 用户手册 / User Manual

## 第一章 / Chapter 1

这是中文内容。This is content in Chinese.

- 列表项 1 / List item 1
- 列表项 2 / List item 2

## 第二章 / Chapter 2

| 表格 | 列 / Column |
|------|------------|
| 数据 / Data | 值 / Value |

Another paragraph with **bold** and *italic* text.

```json
{"name": "test", "value": 42}
```
"""
    md_file = tmp_path / "test.md"
    md_file.write_text(content, encoding="utf-8")
    return md_file


@pytest.fixture
def sample_md_with_table(tmp_path: Path) -> Path:
    """Create a markdown file with tables for XLSX/CSV testing."""
    content = """# Data Table

| Name | Age | City |
|------|-----|------|
| Alice | 30 | Beijing |
| Bob | 25 | Shanghai |
| Carol | 35 | Guangzhou |
"""
    md_file = tmp_path / "table.md"
    md_file.write_text(content, encoding="utf-8")
    return md_file


@pytest.fixture
def sample_md_with_srt(tmp_path: Path) -> Path:
    """Create a markdown file with SRT subtitle blocks."""
    content = """# Subtitles

1
00:00:00,000 --> 00:00:02,500
这是第一条字幕 / This is the first subtitle

2
00:00:03,000 --> 00:00:05,000
这是第二条字幕 / This is the second subtitle

3
00:00:06,000 --> 00:00:08,500
这是第三条字幕 / This is the third subtitle
"""
    md_file = tmp_path / "subtitles.md"
    md_file.write_text(content, encoding="utf-8")
    return md_file


@pytest.fixture
def sample_skeleton_docx(tmp_path: Path) -> Path:
    """Create a minimal skeleton DOCX for XLIFF testing."""
    docx_file = tmp_path / "skeleton.docx"
    with zipfile.ZipFile(docx_file, "w") as zf:
        document_xml = """<?xml version="1.0" encoding="UTF-8" standalone="yes"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
  <w:body>
    <w:p>
      <w:r>
        <w:t>Hello World</w:t>
      </w:r>
    </w:p>
    <w:p>
      <w:r>
        <w:t>Second Paragraph</w:t>
      </w:r>
    </w:p>
  </w:body>
</w:document>"""
        zf.writestr("word/document.xml", document_xml)
        zf.writestr("[Content_Types].xml", "<ContentTypes/>")
        zf.writestr("word/_rels/document.xml.rels", "<Relationships/>")
    return docx_file


@pytest.fixture
def sample_skeleton_pptx(tmp_path: Path) -> Path:
    """Create a minimal skeleton PPTX for XLIFF testing."""
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    # Title slide
    title_slide = prs.slides.add_slide(prs.slide_layouts[0])
    title_shape = title_slide.shapes.title
    title_shape.text = "Presentation Title"
    body_shape = title_slide.placeholders[1]
    body_shape.text = "Hello World"

    # Content slide
    content_slide = prs.slides.add_slide(prs.slide_layouts[1])
    title_shape = content_slide.shapes.title
    title_shape.text = "Chapter 1"
    body_shape = content_slide.placeholders[1]
    body_shape.text = "Second Paragraph"

    pptx_file = tmp_path / "skeleton.pptx"
    prs.save(str(pptx_file))
    return pptx_file


@pytest.fixture
def sample_skeleton_epub(tmp_path: Path) -> Path:
    """Create a minimal skeleton EPUB for XLIFF testing."""
    epub_file = tmp_path / "skeleton.epub"
    import zipfile

    with zipfile.ZipFile(epub_file, "w") as zf:
        container_xml = """<?xml version="1.0" encoding="UTF-8"?>
<container version="1.0" xmlns="urn:oasis:names:tc:opendocument:xmlns:container">
    <rootfiles>
        <rootfile full-path="OEBPS/content.opf" media-type="application/oebps-package+xml"/>
    </rootfiles>
</container>"""
        content_opf = """<?xml version="1.0" encoding="UTF-8"?>
<package xmlns="http://www.idpf.org/2007/opf" version="3.0">
    <metadata/>
    <manifest>
        <item id="chapter1" href="chapter1.xhtml" media-type="application/xhtml+xml"/>
    </manifest>
    <spine>
        <itemref idref="chapter1"/>
    </spine>
</package>"""
        chapter1 = """<?xml version="1.0" encoding="UTF-8"?>
<html xmlns="http://www.w3.org/1999/xhtml">
<head><title>Chapter</title></head>
<body>
<h1>Hello World</h1>
<p>Second Paragraph</p>
</body>
</html>"""
        zf.writestr("META-INF/container.xml", container_xml)
        zf.writestr("OEBPS/content.opf", content_opf)
        zf.writestr("OEBPS/chapter1.xhtml", chapter1)
    return epub_file


@pytest.fixture
def sample_skeleton_html(tmp_path: Path) -> Path:
    """Create a minimal HTML template for XLIFF testing."""
    html_content = """<!DOCTYPE html>
<html>
<head><title>Test Document</title></head>
<body>
<h1>Hello World</h1>
<p>Second Paragraph</p>
</body>
</html>"""
    html_file = tmp_path / "template.html"
    html_file.write_text(html_content, encoding="utf-8")
    return html_file


@pytest.fixture
def sample_skeleton_odf(tmp_path: Path) -> Path:
    """Create a minimal ODF (ODT) skeleton for XLIFF testing."""
    odt_file = tmp_path / "skeleton.odt"
    try:
        from odf.opendocument import OpenDocumentText
        from odf.style import Style, TextProperties
        from odf.text import P, H

        doc = OpenDocumentText()
        h_style = Style(name="Heading1", family="paragraph")
        doc.styles.addElement(h_style)
        doc.text.addElement(H(outlinelevel=1, stylename=h_style, text="Hello World"))
        doc.text.addElement(P(text="Second Paragraph"))
        doc.save(str(odt_file))
    except ImportError:
        # Fallback: create minimal ZIP-based ODT
        with zipfile.ZipFile(odt_file, "w") as zf:
            content_xml = """<?xml version="1.0" encoding="UTF-8"?>
<office:document xmlns:office="urn:oasis:names:tc:opendocument:xmlns:office:document">
  <office:body>
    <office:text>
      <text:p>Hello World</text:p>
    </office:text>
  </office:body>
</office:document>"""
            zf.writestr("content.xml", content_xml)
    return odt_file


@pytest.fixture
def sample_xliff(tmp_path: Path) -> Path:
    """Create a sample XLIFF file for testing."""
    xliff_content = """<?xml version="1.0" encoding="UTF-8"?>
<xliff version="1.2" xmlns="urn:oasis:names:tc:xliff:document:1.2">
  <file original="test.docx" source-language="en" target-language="zh-CN">
    <body>
      <trans-unit id="1">
        <source>Hello World</source>
        <target>你好 世界</target>
      </trans-unit>
      <trans-unit id="2">
        <source>Second Paragraph</source>
        <target>第二段</target>
      </trans-unit>
    </body>
  </file>
</xliff>"""
    xliff_file = tmp_path / "test.xlf"
    xliff_file.write_text(xliff_content, encoding="utf-8")
    return xliff_file


# =============================================================================
# MD → Format Converters (pytest.mark.requires_orf)
# =============================================================================


class TestMD2DOCX:
    """MD → DOCX conversion tests."""

    @pytest.mark.requires_orf
    def test_md2docx_convert(self, sample_md: Path, tmp_path: Path):
        """MD2DOCXConverter should convert MD to DOCX."""
        from orf.channels.md2docx import MD2DOCXConverter

        output = tmp_path / "output.docx"
        converter = MD2DOCXConverter()
        result = converter.convert(sample_md, output)

        assert result.success is True
        assert output.exists()
        assert output.stat().st_size > 0

        # Verify DOCX structure
        with zipfile.ZipFile(output, "r") as zf:
            assert "word/document.xml" in zf.namelist()

    @pytest.mark.requires_orf
    def test_md2docx_output_valid(self, sample_md: Path, tmp_path: Path):
        """MD2DOCX output should be a valid DOCX file."""
        from orf.channels.md2docx import MD2DOCXConverter

        output = tmp_path / "output.docx"
        converter = MD2DOCXConverter()
        result = converter.convert(sample_md, output)

        assert output.exists()
        # Check ZIP validity
        with zipfile.ZipFile(output, "r") as zf:
            namelist = zf.namelist()
            assert "word/document.xml" in namelist
            doc_xml = zf.read("word/document.xml")
            assert b"<w:document" in doc_xml or b"<w:body" in doc_xml


class TestMD2PPTX:
    """MD → PPTX conversion tests."""

    @pytest.mark.requires_orf
    def test_md2pptx_convert(self, sample_md: Path, tmp_path: Path):
        """MD2PPTXConverter should convert MD to PPTX."""
        try:
            from orf.channels.md2pptx import MD2PPTXConverter
        except ImportError:
            pytest.skip("md2pptx not installed")

        output = tmp_path / "output.pptx"
        converter = MD2PPTXConverter()
        result = converter.convert(sample_md, output)

        assert result.success is True
        assert output.exists()
        assert output.stat().st_size > 0

    @pytest.mark.requires_orf
    def test_md2pptx_output_valid(self, sample_md: Path, tmp_path: Path):
        """MD2PPTX output should be a valid PPTX file."""
        try:
            from orf.channels.md2pptx import MD2PPTXConverter
        except ImportError:
            pytest.skip("md2pptx not installed")

        output = tmp_path / "output.pptx"
        converter = MD2PPTXConverter()
        result = converter.convert(sample_md, output)

        assert output.exists()
        with zipfile.ZipFile(output, "r") as zf:
            namelist = zf.namelist()
            assert any("slides/slide" in n for n in namelist)


class TestMD2EPUB:
    """MD → EPUB conversion tests."""

    @pytest.mark.requires_orf
    def test_md2epub_convert(self, sample_md: Path, tmp_path: Path):
        """MD2EPUBConverter should convert MD to EPUB."""
        from orf.channels.md2epub import MD2EPUBConverter

        output = tmp_path / "output.epub"
        converter = MD2EPUBConverter()
        result = converter.convert(sample_md, output)

        assert result.success is True
        assert output.exists()
        assert output.stat().st_size > 0

    @pytest.mark.requires_orf
    def test_md2epub_output_valid(self, sample_md: Path, tmp_path: Path):
        """MD2EPUB output should be a valid EPUB file."""
        from orf.channels.md2epub import MD2EPUBConverter

        output = tmp_path / "output.epub"
        converter = MD2EPUBConverter()
        result = converter.convert(sample_md, output)

        assert output.exists()
        with zipfile.ZipFile(output, "r") as zf:
            namelist = zf.namelist()
            assert "META-INF/container.xml" in namelist


class TestMD2HTML:
    """MD → HTML conversion tests."""

    @pytest.mark.requires_orf
    def test_md2html_convert(self, sample_md: Path, tmp_path: Path):
        """MD2HTMLConverter should convert MD to HTML."""
        from orf.channels.md2html import MD2HTMLConverter

        output = tmp_path / "output.html"
        converter = MD2HTMLConverter()
        result = converter.convert(sample_md, output)

        assert result.success is True
        assert output.exists()
        assert output.stat().st_size > 0

    @pytest.mark.requires_orf
    def test_md2html_output_valid(self, sample_md: Path, tmp_path: Path):
        """MD2HTML output should be valid HTML."""
        from orf.channels.md2html import MD2HTMLConverter

        output = tmp_path / "output.html"
        converter = MD2HTMLConverter()
        result = converter.convert(sample_md, output)

        assert output.exists()
        content = output.read_text(encoding="utf-8")
        assert "<h1" in content or "<p>" in content or "<html" in content.lower()


class TestMD2ODT:
    """MD → ODT conversion tests."""

    @pytest.mark.requires_orf
    def test_md2odt_convert(self, sample_md: Path, tmp_path: Path):
        """MD2ODTConverter should convert MD to ODT."""
        from orf.channels.md2odt import MD2ODTConverter

        output = tmp_path / "output.odt"
        converter = MD2ODTConverter()
        result = converter.convert(sample_md, output)

        assert result.success is True
        assert output.exists()
        assert output.stat().st_size > 0

    @pytest.mark.requires_orf
    def test_md2odt_output_valid(self, sample_md: Path, tmp_path: Path):
        """MD2ODT output should be a valid ODT file."""
        from orf.channels.md2odt import MD2ODTConverter

        output = tmp_path / "output.odt"
        converter = MD2ODTConverter()
        result = converter.convert(sample_md, output)

        assert output.exists()
        with zipfile.ZipFile(output, "r") as zf:
            assert "content.xml" in zf.namelist()


class TestMD2PDF:
    """MD → PDF conversion tests."""

    @pytest.mark.requires_orf
    def test_md2pdf_convert(self, sample_md: Path, tmp_path: Path):
        """MD2PDFConverter should convert MD to PDF."""
        from orf.channels.md2pdf import MD2PDFConverter

        output = tmp_path / "output.pdf"
        converter = MD2PDFConverter()
        result = converter.convert(sample_md, output)

        assert result.success is True
        assert output.exists()
        assert output.stat().st_size > 0

    @pytest.mark.requires_orf
    def test_md2pdf_output_valid(self, sample_md: Path, tmp_path: Path):
        """MD2PDF output should be a valid PDF file."""
        from orf.channels.md2pdf import MD2PDFConverter

        output = tmp_path / "output.pdf"
        converter = MD2PDFConverter()
        result = converter.convert(sample_md, output)

        assert output.exists()
        content = output.read_bytes()
        assert content[:4] == b"%PDF"


class TestMD2RTF:
    """MD → RTF conversion tests."""

    @pytest.mark.requires_orf
    def test_md2rtf_convert(self, sample_md: Path, tmp_path: Path):
        """MD2RTFConverter should convert MD to RTF."""
        from orf.channels.md2rtf import MD2RTFConverter

        output = tmp_path / "output.rtf"
        converter = MD2RTFConverter()
        result = converter.convert(sample_md, output)

        assert result.success is True
        assert output.exists()
        assert output.stat().st_size > 0

    @pytest.mark.requires_orf
    def test_md2rtf_output_valid(self, sample_md: Path, tmp_path: Path):
        """MD2RTF output should be valid RTF."""
        from orf.channels.md2rtf import MD2RTFConverter

        output = tmp_path / "output.rtf"
        converter = MD2RTFConverter()
        result = converter.convert(sample_md, output)

        assert output.exists()
        content = output.read_bytes()
        assert b"{\\rtf" in content[:20] or b"{" in content[:5]


class TestMD2ICML:
    """MD → ICML conversion tests."""

    @pytest.mark.requires_orf
    def test_md2icml_convert(self, sample_md: Path, tmp_path: Path):
        """MD2ICMLConverter should convert MD to ICML."""
        from orf.channels.md2icml import MD2ICMLConverter

        output = tmp_path / "output.icml"
        converter = MD2ICMLConverter()
        result = converter.convert(sample_md, output)

        assert result.success is True
        assert output.exists()
        assert output.stat().st_size > 0

    @pytest.mark.requires_orf
    def test_md2icml_output_valid(self, sample_md: Path, tmp_path: Path):
        """MD2ICML output should be valid XML."""
        from orf.channels.md2icml import MD2ICMLConverter

        output = tmp_path / "output.icml"
        converter = MD2ICMLConverter()
        result = converter.convert(sample_md, output)

        assert output.exists()
        content = output.read_text(encoding="utf-8")
        assert "<xml" in content or "<idPkg" in content


class TestMD2SRT:
    """MD → SRT conversion tests."""

    @pytest.mark.requires_orf
    def test_md2srt_convert(self, sample_md_with_srt: Path, tmp_path: Path):
        """MD2SRTConverter should extract SRT blocks from MD."""
        from orf.channels.md2srt import MD2SRTConverter

        output = tmp_path / "output.srt"
        converter = MD2SRTConverter()
        result = converter.convert(sample_md_with_srt, output)

        assert result.success is True
        assert output.exists()
        assert output.stat().st_size > 0

    @pytest.mark.requires_orf
    def test_md2srt_output_valid(self, sample_md_with_srt: Path, tmp_path: Path):
        """MD2SRT output should be valid SRT format."""
        from orf.channels.md2srt import MD2SRTConverter

        output = tmp_path / "output.srt"
        converter = MD2SRTConverter()
        result = converter.convert(sample_md_with_srt, output)

        assert output.exists()
        content = output.read_text(encoding="utf-8")
        # SRT format: number, timecode, text
        assert "00:00:00,000 --> 00:00:02,500" in content
        assert "1" in content or "2" in content


class TestMD2XLSX:
    """MD → XLSX conversion tests."""

    @pytest.mark.requires_orf
    def test_md2xlsx_convert(self, sample_md_with_table: Path, tmp_path: Path):
        """MD2XLSXConverter should convert MD table to XLSX."""
        from orf.channels.md2xlsx import MD2XLSXConverter

        output = tmp_path / "output.xlsx"
        converter = MD2XLSXConverter()
        result = converter.convert(sample_md_with_table, output)

        assert result.success is True
        assert output.exists()
        assert output.stat().st_size > 0

    @pytest.mark.requires_orf
    def test_md2xlsx_output_valid(self, sample_md_with_table: Path, tmp_path: Path):
        """MD2XLSX output should be a valid XLSX file."""
        from orf.channels.md2xlsx import MD2XLSXConverter

        output = tmp_path / "output.xlsx"
        converter = MD2XLSXConverter()
        result = converter.convert(sample_md_with_table, output)

        assert output.exists()
        with zipfile.ZipFile(output, "r") as zf:
            namelist = zf.namelist()
            assert any("sheet" in n.lower() for n in namelist)


class TestMD2CSV:
    """MD → CSV conversion tests."""

    @pytest.mark.requires_orf
    def test_md2csv_convert(self, sample_md_with_table: Path, tmp_path: Path):
        """MD2CSVConverter should convert MD table to CSV."""
        from orf.channels.md2csv import MD2CSVConverter

        output = tmp_path / "output.csv"
        converter = MD2CSVConverter()
        result = converter.convert(sample_md_with_table, output)

        assert result.success is True
        assert output.exists()
        assert output.stat().st_size > 0

    @pytest.mark.requires_orf
    def test_md2csv_output_valid(self, sample_md_with_table: Path, tmp_path: Path):
        """MD2CSV output should be valid CSV format."""
        from orf.channels.md2csv import MD2CSVConverter

        output = tmp_path / "output.csv"
        converter = MD2CSVConverter()
        result = converter.convert(sample_md_with_table, output)

        assert output.exists()
        content = output.read_text(encoding="utf-8")
        assert "Name" in content or "Age" in content
        assert "," in content


class TestMD2JSON:
    """MD → JSON conversion tests."""

    @pytest.mark.requires_orf
    def test_md2json_convert(self, sample_md: Path, tmp_path: Path):
        """MD2JSONConverter should extract JSON from MD code blocks."""
        from orf.channels.md2json import MD2JSONConverter

        output = tmp_path / "output.json"
        converter = MD2JSONConverter()
        result = converter.convert(sample_md, output)

        assert result.success is True
        assert output.exists()
        assert output.stat().st_size > 0

    @pytest.mark.requires_orf
    def test_md2json_output_valid(self, sample_md: Path, tmp_path: Path):
        """MD2JSON output should be valid JSON."""
        from orf.channels.md2json import MD2JSONConverter

        output = tmp_path / "output.json"
        converter = MD2JSONConverter()
        result = converter.convert(sample_md, output)

        assert output.exists()
        content = output.read_text(encoding="utf-8")
        data = json.loads(content)
        assert data["name"] == "test"
        assert data["value"] == 42


# =============================================================================
# XLIFF → Format Converters
# =============================================================================


class TestXLIFF2DOCX:
    """XLIFF → DOCX conversion tests."""

    @pytest.mark.requires_orf
    def test_xliff2docx_convert(
        self,
        sample_skeleton_docx: Path,
        sample_xliff: Path,
        tmp_path: Path,
    ):
        """XLIFF2DOCXConverter should backfill XLIFF into DOCX skeleton."""
        from orf.channels.xliff2docx import XLIFF2DOCXConverter

        output = tmp_path / "output.docx"
        converter = XLIFF2DOCXConverter()
        result = converter.convert(sample_skeleton_docx, sample_xliff, output)

        assert result.success is True
        assert output.exists()
        assert output.stat().st_size > 0

    @pytest.mark.requires_orf
    def test_xliff2docx_output_valid(
        self,
        sample_skeleton_docx: Path,
        sample_xliff: Path,
        tmp_path: Path,
    ):
        """XLIFF2DOCX output should be a valid DOCX file."""
        from orf.channels.xliff2docx import XLIFF2DOCXConverter

        output = tmp_path / "output.docx"
        converter = XLIFF2DOCXConverter()
        result = converter.convert(sample_skeleton_docx, sample_xliff, output)

        assert output.exists()
        with zipfile.ZipFile(output, "r") as zf:
            assert "word/document.xml" in zf.namelist()


class TestXLIFF2PPTX:
    """XLIFF → PPTX conversion tests."""

    @pytest.mark.requires_orf
    def test_xliff2pptx_convert(
        self,
        sample_skeleton_pptx: Path,
        sample_xliff: Path,
        tmp_path: Path,
    ):
        """XLIFF2PPTXConverter should backfill XLIFF into PPTX skeleton."""
        from orf.channels.xliff2pptx import XLIFF2PPTXConverter

        output = tmp_path / "output.pptx"
        converter = XLIFF2PPTXConverter()
        result = converter.convert(sample_skeleton_pptx, sample_xliff, output)

        assert result.success is True
        assert output.exists()
        assert output.stat().st_size > 0

    @pytest.mark.requires_orf
    def test_xliff2pptx_output_valid(
        self,
        sample_skeleton_pptx: Path,
        sample_xliff: Path,
        tmp_path: Path,
    ):
        """XLIFF2PPTX output should be a valid PPTX file."""
        from orf.channels.xliff2pptx import XLIFF2PPTXConverter

        output = tmp_path / "output.pptx"
        converter = XLIFF2PPTXConverter()
        result = converter.convert(sample_skeleton_pptx, sample_xliff, output)

        assert output.exists()
        with zipfile.ZipFile(output, "r") as zf:
            namelist = zf.namelist()
            assert any("slides/slide" in n for n in namelist)


class TestXLIFF2EPUB:
    """XLIFF → EPUB conversion tests."""

    @pytest.mark.requires_orf
    def test_xliff2epub_convert(
        self,
        sample_skeleton_epub: Path,
        sample_xliff: Path,
        tmp_path: Path,
    ):
        """XLIFF2EPUBConverter should backfill XLIFF into EPUB skeleton."""
        from orf.channels.xliff2epub import XLIFF2EPUBConverter

        output = tmp_path / "output.epub"
        converter = XLIFF2EPUBConverter()
        result = converter.convert(sample_skeleton_epub, sample_xliff, output)

        assert result.success is True
        assert output.exists()
        assert output.stat().st_size > 0

    @pytest.mark.requires_orf
    def test_xliff2epub_output_valid(
        self,
        sample_skeleton_epub: Path,
        sample_xliff: Path,
        tmp_path: Path,
    ):
        """XLIFF2EPUB output should be a valid EPUB file."""
        from orf.channels.xliff2epub import XLIFF2EPUBConverter

        output = tmp_path / "output.epub"
        converter = XLIFF2EPUBConverter()
        result = converter.convert(sample_skeleton_epub, sample_xliff, output)

        assert output.exists()
        with zipfile.ZipFile(output, "r") as zf:
            namelist = zf.namelist()
            assert "META-INF/container.xml" in namelist


class TestXLIFF2HTML:
    """XLIFF → HTML conversion tests."""

    @pytest.mark.requires_orf
    def test_xliff2html_convert(
        self,
        sample_skeleton_html: Path,
        sample_xliff: Path,
        tmp_path: Path,
    ):
        """XLIFF2HTMLConverter should backfill XLIFF into HTML template."""
        from orf.channels.xliff2html import XLIFF2HTMLConverter

        output = tmp_path / "output.html"
        converter = XLIFF2HTMLConverter()
        result = converter.convert(sample_skeleton_html, sample_xliff, output)

        assert result.success is True
        assert output.exists()
        assert output.stat().st_size > 0

    @pytest.mark.requires_orf
    def test_xliff2html_output_valid(
        self,
        sample_skeleton_html: Path,
        sample_xliff: Path,
        tmp_path: Path,
    ):
        """XLIFF2HTML output should be valid HTML."""
        from orf.channels.xliff2html import XLIFF2HTMLConverter

        output = tmp_path / "output.html"
        converter = XLIFF2HTMLConverter()
        result = converter.convert(sample_skeleton_html, sample_xliff, output)

        assert output.exists()
        content = output.read_text(encoding="utf-8")
        assert "<html" in content.lower() or "<body" in content.lower()


class TestXLIFF2ODF:
    """XLIFF → ODF conversion tests."""

    @pytest.mark.requires_orf
    def test_xliff2odf_convert(
        self,
        sample_skeleton_odf: Path,
        sample_xliff: Path,
        tmp_path: Path,
    ):
        """XLIFF2ODFConverter should backfill XLIFF into ODF skeleton."""
        try:
            from orf.channels.xliff2odf import XLIFF2ODFConverter
        except ImportError:
            pytest.skip("translate-toolkit not installed")

        output = tmp_path / "output.odt"
        converter = XLIFF2ODFConverter()
        result = converter.convert(sample_skeleton_odf, sample_xliff, output)

        assert result.success is True
        assert output.exists()
        assert output.stat().st_size > 0

    @pytest.mark.requires_orf
    def test_xliff2odf_output_valid(
        self,
        sample_skeleton_odf: Path,
        sample_xliff: Path,
        tmp_path: Path,
    ):
        """XLIFF2ODF output should be a valid ODF file."""
        try:
            from orf.channels.xliff2odf import XLIFF2ODFConverter
        except ImportError:
            pytest.skip("translate-toolkit not installed")

        output = tmp_path / "output.odt"
        converter = XLIFF2ODFConverter()
        result = converter.convert(sample_skeleton_odf, sample_xliff, output)

        assert output.exists()
        with zipfile.ZipFile(output, "r") as zf:
            assert "content.xml" in zf.namelist()


# =============================================================================
# Validation Helpers
# =============================================================================


def validate_docx(path: Path) -> tuple[bool, list[str]]:
    """Validate DOCX file structure. Returns (is_valid, errors)."""
    errors = []
    if not path.exists():
        return False, ["File does not exist"]
    if path.stat().st_size == 0:
        return False, ["File is empty"]
    try:
        with zipfile.ZipFile(path, "r") as zf:
            if "word/document.xml" not in zf.namelist():
                errors.append("Missing word/document.xml")
    except zipfile.BadZipFile:
        return False, ["Not a valid ZIP/DOCX"]
    return len(errors) == 0, errors


def validate_pptx(path: Path) -> tuple[bool, list[str]]:
    """Validate PPTX file structure. Returns (is_valid, errors)."""
    errors = []
    if not path.exists():
        return False, ["File does not exist"]
    if path.stat().st_size == 0:
        return False, ["File is empty"]
    try:
        with zipfile.ZipFile(path, "r") as zf:
            namelist = zf.namelist()
            if not any("slides/slide" in n for n in namelist):
                errors.append("No slides found")
    except zipfile.BadZipFile:
        return False, ["Not a valid ZIP/PPTX"]
    return len(errors) == 0, errors


def validate_epub(path: Path) -> tuple[bool, list[str]]:
    """Validate EPUB file structure. Returns (is_valid, errors)."""
    errors = []
    if not path.exists():
        return False, ["File does not exist"]
    if path.stat().st_size == 0:
        return False, ["File is empty"]
    try:
        with zipfile.ZipFile(path, "r") as zf:
            namelist = zf.namelist()
            if "META-INF/container.xml" not in namelist:
                errors.append("Missing META-INF/container.xml")
    except zipfile.BadZipFile:
        return False, ["Not a valid ZIP/EPUB"]
    return len(errors) == 0, errors


def validate_pdf(path: Path) -> tuple[bool, list[str]]:
    """Validate PDF file. Returns (is_valid, errors)."""
    errors = []
    if not path.exists():
        return False, ["File does not exist"]
    if path.stat().st_size == 0:
        return False, ["File is empty"]
    header = path.read_bytes()[:4]
    if header != b"%PDF":
        errors.append(f"Invalid PDF header: {header}")
    return len(errors) == 0, errors


def validate_xlsx(path: Path) -> tuple[bool, list[str]]:
    """Validate XLSX file structure. Returns (is_valid, errors)."""
    errors = []
    if not path.exists():
        return False, ["File does not exist"]
    if path.stat().st_size == 0:
        return False, ["File is empty"]
    try:
        with zipfile.ZipFile(path, "r") as zf:
            namelist = zf.namelist()
            if not any("sheet" in n.lower() for n in namelist):
                errors.append("No sheet found in XLSX")
    except zipfile.BadZipFile:
        return False, ["Not a valid ZIP/XLSX"]
    return len(errors) == 0, errors