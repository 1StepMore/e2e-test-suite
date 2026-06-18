"""E2E tests for ORF MCP server tools.

Tests the ORF MCP server tools:
- apply_md tool
- apply_xliff tool
- batch_convert tool

Each test directly calls the MCP tool functions.
"""

import json
import os
import sys
import tempfile
from pathlib import Path
from unittest.mock import MagicMock, patch

# The ORF MCP path validator caches a global instance built from
# ORF_MCP_ALLOWED_DIRS at module import time. Without setting this,
# tmp_path (which lives under /tmp/pytest-of-...) gets rejected
# before our patch on PathValidator.validate can run. Allow /tmp
# so pytest's per-test temp dirs are accepted.
os.environ.setdefault("ORF_MCP_ALLOWED_DIRS", "/tmp")

import pytest


class TestORFMCP:
    """ORF MCP server tool tests."""

    @pytest.mark.requires_orf
    def test_apply_md_to_docx(self, tmp_path):
        """Test apply_md converts MD to DOCX."""
        md_content = """---
source_lang: en
target_lang: zh
original_file: test.md
processor: "OL"
version: "0.2.6"
translated_at: 2026-05-28T00:00:00Z
---

# Test Document

This is test content for conversion.
"""
        input_md = tmp_path / "test.md"
        input_md.write_text(md_content, encoding="utf-8")

        with patch("orf.mcp.server.PathValidator.validate") as mock_validate, \
             patch("orf.mcp.server._run_cli_command") as mock_cli:

            mock_validate.return_value = (True, None)
            mock_cli.return_value = {
                "success": True,
                "output_path": str(tmp_path / "test.docx"),
                "errors": [],
                "warnings": [],
                "metadata": {},
            }

            from orf.mcp.server import apply_md

            result = apply_md(
                input_md=str(input_md),
                target_format="docx",
                output_path=str(tmp_path / "output.docx"),
            )

            parsed = json.loads(result)
            assert parsed["success"] is True

    @pytest.mark.requires_orf
    def test_apply_md_to_html(self, tmp_path):
        """Test apply_md converts MD to HTML."""
        md_content = """---
source_lang: en
target_lang: zh
---

# HTML Test

HTML content here.
"""
        input_md = tmp_path / "html.md"
        input_md.write_text(md_content, encoding="utf-8")

        with patch("orf.mcp.server.PathValidator.validate") as mock_validate, \
             patch("orf.mcp.server._run_cli_command") as mock_cli:

            mock_validate.return_value = (True, None)
            mock_cli.return_value = {
                "success": True,
                "output_path": str(tmp_path / "html_output.html"),
                "errors": [],
                "warnings": [],
                "metadata": {},
            }

            from orf.mcp.server import apply_md

            result = apply_md(
                input_md=str(input_md),
                target_format="html",
            )

            parsed = json.loads(result)
            assert parsed["success"] is True

    @pytest.mark.requires_orf
    def test_apply_md_invalid_path(self, tmp_path):
        """Test apply_md handles invalid paths."""
        with patch("orf.mcp.server.PathValidator.validate") as mock_validate:
            mock_validate.return_value = (False, "Path not allowed: /invalid/path.md")

            from orf.mcp.server import apply_md

            result = apply_md(
                input_md="/invalid/path.md",
                target_format="docx",
            )

            parsed = json.loads(result)
            assert parsed["success"] is False
            assert "error" in parsed or "errors" in parsed

    @pytest.mark.requires_orf
    def test_apply_xliff_basic(self, tmp_path):
        """Test apply_xliff applies XLIFF to original document."""
        from docx import Document

        doc = Document()
        doc.add_heading("Original", level=1)
        doc.add_paragraph("Content.")

        input_docx = tmp_path / "original.docx"
        doc.save(str(input_docx))

        xliff_content = """<?xml version="1.0" encoding="UTF-8"?>
<xliff version="1.2" xmlns="urn:oasis:names:tc:xliff:document:1.2">
  <file source-language="en" target-language="zh" original="original.docx">
    <body>
      <trans-unit id="1">
        <source>Original</source>
        <target>翻译原文</target>
      </trans-unit>
      <trans-unit id="2">
        <source>Content.</source>
        <target>内容。</target>
      </trans-unit>
    </body>
  </file>
</xliff>
"""
        xliff_path = tmp_path / "translated.xlf"
        xliff_path.write_text(xliff_content, encoding="utf-8")

        with patch("orf.mcp.server.PathValidator.validate") as mock_validate, \
             patch("orf.mcp.server._run_cli_command") as mock_cli:

            def validate_side_effect(path):
                return (True, None)

            mock_validate.side_effect = validate_side_effect
            mock_cli.return_value = {
                "success": True,
                "output_path": str(tmp_path / "result.docx"),
                "errors": [],
                "warnings": [],
                "metadata": {},
            }

            from orf.mcp.server import apply_xliff

            result = apply_xliff(
                input_file=str(input_docx),
                xliff_path=str(xliff_path),
                output_path=str(tmp_path / "output.docx"),
                format="docx",
            )

            parsed = json.loads(result)
            assert parsed["success"] is True

    @pytest.mark.requires_orf
    def test_apply_xliff_with_images(self, tmp_path):
        """Test apply_xliff handles image placement data."""
        from docx import Document

        doc = Document()
        doc.add_heading("Image Test", level=1)

        input_docx = tmp_path / "image_test.docx"
        doc.save(str(input_docx))

        xliff_path = tmp_path / "image.xlf"
        xliff_path.write_text("""<?xml version="1.0"?>
<xliff version="1.2">
  <file source-language="en" target-language="zh">
    <body>
      <trans-unit id="1">
        <source>Image Test</source>
        <target>图片测试</target>
      </trans-unit>
    </body>
  </file>
</xliff>
""", encoding="utf-8")

        images = [
            {
                "data_base64": "iVBORw0KGgoAAAANSUhEUgAAAAEAAAABCAYAAAAfFcSJAAAADUlEQVR42mNk+M9QDwADhgGAWjR9awAAAABJRU5ErkJggg==",
                "mime_type": "image/png",
                "width": 100,
                "height": 100,
                "paragraph_index": 0,
            }
        ]

        with patch("orf.mcp.server.PathValidator.validate") as mock_validate, \
             patch("orf.mcp.server._run_cli_command") as mock_cli:

            mock_validate.return_value = (True, None)
            mock_cli.return_value = {
                "success": True,
                "output_path": str(tmp_path / "with_images.docx"),
                "errors": [],
                "warnings": [],
                "metadata": {},
            }

            from orf.mcp.server import apply_xliff

            result = apply_xliff(
                input_file=str(input_docx),
                xliff_path=str(xliff_path),
                output_path=str(tmp_path / "output.docx"),
                format="docx",
                images=images,
            )

            parsed = json.loads(result)
            assert parsed["success"] is True

    @pytest.mark.requires_orf
    def test_apply_xliff_invalid_input(self, tmp_path):
        """Test apply_xliff handles invalid input paths."""
        with patch("orf.mcp.server.PathValidator.validate") as mock_validate:
            mock_validate.return_value = (False, "Path not allowed")

            from orf.mcp.server import apply_xliff

            result = apply_xliff(
                input_file="/invalid/input.docx",
                xliff_path="/invalid/translated.xlf",
                output_path="/invalid/output.docx",
                format="docx",
            )

            parsed = json.loads(result)
            assert parsed["success"] is False

    @pytest.mark.requires_orf
    def test_batch_convert_basic(self, tmp_path):
        """Test batch_convert processes multiple MD files."""
        # Create MD files with frontmatter
        for i in range(3):
            md_content = f"""---
source_lang: en
target_lang: zh
original_file: batch_{i}.md
processor: "OL"
---

# Batch {i}

Content {i}.
"""
            md_file = tmp_path / f"batch_{i}.md"
            md_file.write_text(md_content, encoding="utf-8")

        with patch("orf.mcp.server.PathValidator.validate") as mock_validate, \
             patch("orf.mcp.server._run_cli_command") as mock_cli:

            mock_validate.return_value = (True, None)
            mock_cli.return_value = {
                "success_count": 3,
                "fail_count": 0,
                "total": 3,
            }

            from orf.mcp.server import batch_convert

            result = batch_convert(
                input_dir=str(tmp_path),
                target_format="docx",
                pattern="batch_*.md",
            )

            parsed = json.loads(result)
            assert "success_count" in parsed or "total" in parsed

    @pytest.mark.requires_orf
    def test_batch_convert_invalid_dir(self, tmp_path):
        """Test batch_convert handles invalid directory."""
        with patch("orf.mcp.server.PathValidator.validate") as mock_validate:
            mock_validate.return_value = (False, "Directory not allowed")

            from orf.mcp.server import batch_convert

            result = batch_convert(
                input_dir="/invalid/directory",
                target_format="docx",
            )

            parsed = json.loads(result)
            assert parsed["success_count"] == 0 or "error" in parsed

    @pytest.mark.requires_orf
    def test_apply_md_to_epub(self, tmp_path):
        """Test apply_md converts MD to EPUB format."""
        md_content = """---
source_lang: en
target_lang: zh
---

# EPUB Test

EPUB content for testing.
"""
        input_md = tmp_path / "epub.md"
        input_md.write_text(md_content, encoding="utf-8")

        with patch("orf.mcp.server.PathValidator.validate") as mock_validate, \
             patch("orf.mcp.server._run_cli_command") as mock_cli:

            mock_validate.return_value = (True, None)
            mock_cli.return_value = {
                "success": True,
                "output_path": str(tmp_path / "epub_output.epub"),
                "errors": [],
                "warnings": [],
                "metadata": {},
            }

            from orf.mcp.server import apply_md

            result = apply_md(
                input_md=str(input_md),
                target_format="epub",
            )

            parsed = json.loads(result)
            assert parsed["success"] is True

    @pytest.mark.requires_orf
    def test_batch_convert_epub(self, tmp_path):
        """Test batch_convert to EPUB format."""
        for i in range(2):
            md_content = f"""---
source_lang: en
target_lang: zh
---

# Batch EPUB {i}

Content {i}.
"""
            md_file = tmp_path / f"epub_batch_{i}.md"
            md_file.write_text(md_content, encoding="utf-8")

        with patch("orf.mcp.server.PathValidator.validate") as mock_validate, \
             patch("orf.mcp.server._run_cli_command") as mock_cli:

            mock_validate.return_value = (True, None)
            mock_cli.return_value = {
                "success_count": 2,
                "fail_count": 0,
                "total": 2,
            }

            from orf.mcp.server import batch_convert

            result = batch_convert(
                input_dir=str(tmp_path),
                target_format="epub",
                pattern="epub_batch_*.md",
            )

            parsed = json.loads(result)
            assert "success_count" in parsed or "total" in parsed

    @pytest.mark.requires_orf
    def test_detect_format_tool(self, tmp_path):
        """Test detect_format tool identifies document format."""
        from docx import Document

        doc = Document()
        doc.add_heading("Detect Test", level=1)

        docx_path = tmp_path / "detect.docx"
        doc.save(str(docx_path))

        with patch("orf.mcp.server.PathValidator.validate") as mock_validate, \
             patch("orf.mcp.server._run_cli_command") as mock_cli:

            mock_validate.return_value = (True, None)
            mock_cli.return_value = {
                "format": "DOCX",
                "size_mb": 0.01,
                "resource_count": 0,
                "manifest_status": "not found",
            }

            from orf.mcp.server import detect_format

            result = detect_format(file_path=str(docx_path))
            parsed = json.loads(result)

            assert parsed["format"] == "DOCX"

    @pytest.mark.requires_orf
    def test_info_tool(self, tmp_path):
        """Test info tool returns document information."""
        from docx import Document

        doc = Document()
        doc.add_heading("Info Test", level=1)
        doc.add_paragraph("Content.")

        docx_path = tmp_path / "info.docx"
        doc.save(str(docx_path))

        with patch("orf.mcp.server.PathValidator.validate") as mock_validate, \
             patch("orf.mcp.server._run_cli_command") as mock_cli:

            mock_validate.return_value = (True, None)
            mock_cli.return_value = {
                "format": "DOCX",
                "size_mb": 0.01,
                "resource_count": 0,
                "manifest_status": "not found",
            }

            from orf.mcp.server import info

            result = info(file_path=str(docx_path))
            parsed = json.loads(result)

            assert parsed["format"] == "DOCX"
            assert "size_mb" in parsed