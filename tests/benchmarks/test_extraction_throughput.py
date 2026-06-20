"""Throughput benchmarks for OPP extraction.

Measures document extraction timing across sample formats.
Marked with @pytest.mark.benchmark — run with: pytest tests/benchmarks/ -m benchmark
"""

import time
from pathlib import Path

import pytest


pytestmark = [
    pytest.mark.benchmark,
    pytest.mark.requires_opp,
]


@pytest.fixture
def opp_pipeline(tmp_path: Path):
    """Create an OPP pipeline instance for benchmarking."""
    from opp.pipeline import OPPPipeline

    resource_dir = tmp_path / "opp_resources"
    resource_dir.mkdir(exist_ok=True)

    return OPPPipeline(resource_storage_dir=resource_dir)


def _create_sample_docx(path: Path, paragraphs: int = 50) -> Path:
    """Create a DOCX with a given number of paragraphs."""
    from docx import Document

    doc = Document()
    doc.add_heading("Benchmark Document", level=1)
    for i in range(paragraphs):
        doc.add_paragraph(
            f"This is paragraph {i} with some content to exercise the extraction "
            f"pipeline. Lorem ipsum dolor sit amet, consectetur adipiscing elit."
        )
    doc.add_heading("Data Table", level=2)
    table = doc.add_table(rows=10, cols=4)
    table.style = "Light Grid Accent 1"
    for r in range(10):
        for c in range(4):
            table.rows[r].cells[c].text = f"R{r}C{c}"
    doc.save(str(path))
    return path


@pytest.fixture
def sample_docx_50(tmp_path: Path) -> Path:
    """50-paragraph DOCX."""
    return _create_sample_docx(tmp_path / "doc_50.docx", paragraphs=50)


@pytest.fixture
def sample_docx_200(tmp_path: Path) -> Path:
    """200-paragraph DOCX for heavier load."""
    return _create_sample_docx(tmp_path / "doc_200.docx", paragraphs=200)


class TestDOCXExtractionThroughput:
    """OPP DOCX extraction timing benchmarks."""

    def test_docx_50_paragraphs(self, opp_pipeline, sample_docx_50, tmp_path):
        """Time extraction of a 50-paragraph DOCX."""
        output_dir = tmp_path / "out_50"
        output_dir.mkdir(exist_ok=True)

        start = time.perf_counter()
        result = opp_pipeline.process_file(sample_docx_50)
        elapsed = time.perf_counter() - start

        assert result is not None
        assert result.extraction_result is not None

        print(f"\n  DOCX (50 paragraphs): {elapsed:.4f}s")
        assert elapsed < 30.0, f"Extraction too slow: {elapsed:.2f}s"

    def test_docx_200_paragraphs(self, opp_pipeline, sample_docx_200, tmp_path):
        """Time extraction of a 200-paragraph DOCX."""
        output_dir = tmp_path / "out_200"
        output_dir.mkdir(exist_ok=True)

        start = time.perf_counter()
        result = opp_pipeline.process_file(sample_docx_200)
        elapsed = time.perf_counter() - start

        assert result is not None
        assert result.extraction_result is not None

        print(f"\n  DOCX (200 paragraphs): {elapsed:.4f}s")
        assert elapsed < 60.0, f"Extraction too slow: {elapsed:.2f}s"

    def test_docx_consecutive_runs(self, opp_pipeline, sample_docx_50, tmp_path):
        """Run extraction 3 times and report min/avg/max."""
        timings = []
        for i in range(3):
            start = time.perf_counter()
            result = opp_pipeline.process_file(sample_docx_50)
            elapsed = time.perf_counter() - start
            timings.append(elapsed)
            assert result is not None

        avg = sum(timings) / len(timings)
        print(f"\n  DOCX (3 runs): min={min(timings):.4f}s  avg={avg:.4f}s  max={max(timings):.4f}s")
        assert avg < 30.0, f"Average extraction time {avg:.2f}s too high"
