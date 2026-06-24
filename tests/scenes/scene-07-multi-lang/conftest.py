"""Fixtures for scene-07 multi-language E2E tests.

Generates synthetic input files for 3 language pairs × 4 formats = 12 cells.
Uses ``indirect=True`` parametrization so the fixture receives ``(lang_pair, fmt)``
as ``request.param``.
"""

from __future__ import annotations

import json
import random
from pathlib import Path

import fitz
import pytest

random.seed(42)

EN_CONTENT = (Path(__file__).parent / "fixtures" / "en_content.md").read_text()


def _make_docx(path: Path, content: str) -> None:
    """Generate a minimal .docx with headings and paragraphs from markdown text."""
    from docx import Document

    doc = Document()
    for line in content.split("\n"):
        if line.startswith("# "):
            doc.add_heading(line[2:], level=1)
        elif line.strip():
            doc.add_paragraph(line)
    doc.save(str(path))


def _make_pdf(path: Path, content: str) -> None:
    """Generate a minimal .pdf with text lines via PyMuPDF (fitz)."""
    doc = fitz.open()
    page = doc.new_page()
    y = 50
    for line in content.split("\n"):
        if line.strip():
            page.insert_text((50, y), line)
            y += 20
    doc.save(str(path))
    doc.close()


@pytest.fixture(scope="session")
def sample_input(request, tmp_path_factory):
    """Generate an input file for one (lang_pair, fmt) cell.

    Uses ``indirect=True`` parametrization — ``request.param`` is a
    ``(lang_pair, fmt)`` tuple, e.g. ``("en-fr", "docx")``.

    Returns ``(path: Path, lang_pair: str, fmt: str)``.
    """
    lang_pair, fmt = request.param
    tmp = tmp_path_factory.mktemp(f"fixtures_{lang_pair}_{fmt}")

    if fmt == "docx":
        path = tmp / "sample.docx"
        _make_docx(path, EN_CONTENT)
    elif fmt == "pdf":
        path = tmp / "sample.pdf"
        _make_pdf(path, EN_CONTENT)
    elif fmt == "html":
        path = tmp / "sample.html"
        html_body = EN_CONTENT.replace("\n", "<br>\n")
        path.write_text(
            (
                "<!DOCTYPE html>\n<html><head><meta charset=\"utf-8\"></head>\n"
                f"<body><p>{html_body}</p></body>\n</html>\n"
            ),
            encoding="utf-8",
        )
    elif fmt == "json":
        path = tmp / "sample.json"
        data = {
            "company": "TechCorp",
            "name": "Alice Johnson",
            "email": "alice@techcorp.com",
            "phone": "+1-555-1234",
        }
        path.write_text(json.dumps(data, indent=2, ensure_ascii=False), encoding="utf-8")
    else:
        raise ValueError(f"Unsupported format: {fmt}")

    return (path, lang_pair, fmt)
